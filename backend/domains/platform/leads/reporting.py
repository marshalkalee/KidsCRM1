"""Выгрузка заявок в Excel (TRU-95): что видно в таблице с фильтрами."""

import io

import openpyxl
from django.utils import timezone
from openpyxl.styles import Font

COLUMNS = [
    ("Ребёнок", 22),
    ("Возраст", 9),
    ("Родитель", 22),
    ("Телефон", 16),
    ("Направление", 18),
    ("Источник", 16),
    ("Статус", 20),
    ("Причина отказа", 20),
    ("Ответственный", 22),
    ("Филиал", 18),
    ("Создана", 17),
    ("Дней в статусе", 14),
]


def leads_workbook(leads, funnel) -> bytes:
    """funnel — этапы центра (stages.Funnel): в колонке «Статус» — его названия."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Заявки"
    sheet.append([title for title, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    now = timezone.now()
    for lead in leads:
        sheet.append(
            [
                lead.child_name,
                lead.child_age,
                lead.parent_name,
                lead.phone,
                lead.direction.name if lead.direction else "",
                lead.source.name if lead.source else "",
                funnel.of(lead).name,
                lead.rejection_reason.name if lead.rejection_reason else "",
                lead.assigned_to.full_name if lead.assigned_to else "",
                lead.branch.name if lead.branch else "",
                # Excel не понимает дату с часовым поясом — местное время центра.
                timezone.localtime(lead.created_at).replace(tzinfo=None),
                max(0, (now - lead.status_changed_at).days),
            ]
        )
        sheet.cell(row=sheet.max_row, column=11).number_format = "DD.MM.YYYY HH:MM"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
