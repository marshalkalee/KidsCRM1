"""Главный экран дашборда владельца и приёмка M3 (TRU-129): цифры дашборда
сверены с операционными экранами и подробными отчётами, права ролей — на
всех отчётах M3, выгрузка — на каждом отчёте."""

import io
from datetime import datetime, time, timedelta
from unittest import mock

import openpyxl
from django.test import tag

from domains.platform.leads.models import Lead, LeadStatusChange
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.tenants.org_settings import ADMIN_SEES_ORG_SUMMARY
from domains.platform.users.models import User
from domains.scheduling.groups.models import Group, GroupMembership

from .reports import REPORTS
from .tests_forecast import ForecastFixtures

URL = "/api/v1/analytics/dashboard/"
API = "/api/v1/analytics/"

# Все отчёты M3 на чтение: путь и параметры, без которых отчёт ответит 400.
ANALYTICS_ENDPOINTS = [
    ("dashboard/", {}),
    ("metrics/", {"metrics": "revenue"}),
    ("breakdown/", {"metric": "revenue", "by": "method"}),
    ("heatmap/", {}),
    ("attendance-trends/", {}),
    ("risk-list/", {}),
    ("group-occupancy/", {}),
    ("teacher-workload/", {}),
    ("funnel/", {}),
    ("funnel/by/", {"by": "source"}),
    ("forecast/", {}),
    ("renewal-conversion/", {}),
    ("churn/", {}),
    ("rejections/", {}),
    ("rejections/by/", {"by": "source"}),
    ("branches/", {}),
    ("branches/trend/", {"metric": "revenue"}),
    ("sources/", {}),
    ("export/", {"report": "revenue"}),
    ("catalog/", {}),
]


class DashboardFixtures(ForecastFixtures):
    def setUp(self):
        super().setUp()
        self.n = 0
        # Выручка: 30 000 в Абая, 15 000 в Саина; в Саина недоплата 10 000.
        self.pay(self.subscription(self.abaya), 30000)
        self.pay(self.subscription(self.saina, price=25000), 15000)
        # Продления: 25 закончились давно, 10 продлены — прогноз «диапазон».
        self.ended(25, renewed=10, ends_on=self.history_end())
        # Группы: 3 из 10 и 8 из 10 в Абая, 2 из 4 в Саина.
        self.group(self.abaya, capacity=10, members=3)
        self.group(self.abaya, capacity=10, members=8)
        self.group(self.saina, capacity=4, members=2)
        # Заявки: 4 новые, из них 1 купила.
        for path in (["new"], ["new", "contacted"], ["new", "purchased"], ["new"]):
            self.lead(*path)
        self.client.force_authenticate(self.owner)

    def group(self, branch, *, capacity, members):
        self.n += 1
        group = Group.objects.create(
            organization=self.org,
            branch=branch,
            direction=self.ballet,
            name=f"Группа {self.n}",
            capacity=capacity,
        )
        for _ in range(members):
            GroupMembership.objects.create(
                organization=self.org,
                group=group,
                child=self.child(),
                joined_at=self.today - timedelta(days=10),
            )
        return group

    def lead(self, *path, branch=None):
        self.n += 1
        lead = Lead.objects.create(
            organization=self.org,
            kind="new",
            parent_name=f"Родитель {self.n}",
            phone=f"+7702{self.n:07d}",
            branch=branch or self.abaya,
            direction=self.ballet,
            status=path[-1],
        )
        created = self.tz.localize(datetime.combine(self.today, time(9)))
        Lead.objects.filter(pk=lead.pk).update(created_at=created)
        previous = ""
        for status in path:
            LeadStatusChange.objects.create(
                organization=self.org, lead=lead, from_status=previous, to_status=status
            )
            previous = status
        return lead

    def period(self):
        return {
            "period": "custom",
            "from": (self.history_end() - timedelta(days=10)).isoformat(),
            "to": self.today.isoformat(),
        }

    def get(self, path, **params):
        response = self.client.get(f"{API}{path}", {**self.period(), **params})
        self.assertEqual(response.status_code, 200, response.content[:300])
        return response.data

    def tiles(self, **params):
        return self.get("dashboard/", **params)["tiles"]


class CrossCheckTests(DashboardFixtures):
    """Сквозная сверка: каждая плитка = цифра своего экрана."""

    def test_revenue_is_the_metrics_api(self):
        tile = self.tiles()["revenue"]
        metric = self.get("metrics/", metrics="revenue", series="0")["metrics"]["revenue"]
        self.assertEqual(tile["value"], metric["value"])
        self.assertEqual(tile["value"], "45000")
        self.assertEqual(tile["change_percent"], metric["change_percent"])

    def test_debt_is_the_debtors_screen(self):
        screen = self.client.get("/api/v1/subscriptions/debtors/").data
        tile = self.tiles()["debt"]
        self.assertEqual(tile["value"], screen["total_debt"])
        # Филиал в шапке сужает обе цифры одинаково.
        header = {"HTTP_X_BRANCH_ID": str(self.saina.pk)}
        screen = self.client.get("/api/v1/subscriptions/debtors/", **header).data
        tile = self.client.get(URL, self.period(), **header).data["tiles"]["debt"]
        self.assertEqual(tile["value"], screen["total_debt"])
        self.assertEqual(tile["value"], "10000")

    def test_group_fill_is_the_groups_list(self):
        groups = self.client.get("/api/v1/groups/", {"status": "active"}).data["results"]
        tile = self.tiles()["group_fill"]
        members = sum(g["members_count"] for g in groups)
        capacity = sum(g["capacity"] for g in groups)
        self.assertEqual((tile["members"], tile["capacity"]), (members, capacity))
        self.assertEqual((members, capacity), (13, 24))
        self.assertEqual(float(tile["value"]), round(members * 100 / capacity, 1))
        self.assertEqual(tile["underfilled"], sum(g["is_underfilled"] for g in groups))
        self.assertEqual(tile["groups"], 3)

    def test_lead_conversion_is_the_funnel(self):
        funnel = self.get("funnel/")["funnel"]
        tile = self.tiles()["lead_conversion"]
        self.assertEqual(tile["value"], funnel["conversion"])
        self.assertEqual((tile["leads"], tile["purchased"], tile["value"]), (4, 1, 25.0))

    def test_renewal_conversion_is_the_renewals_report(self):
        summary = self.get("renewal-conversion/")["summary"]
        tile = self.tiles()["renewal_conversion"]
        self.assertEqual(tile["value"], summary["rate"])
        self.assertEqual((tile["decided"], tile["renewed"]), (summary["decided"], 10))
        self.assertEqual(tile["value"], 40.0)

    def test_risk_is_the_risk_list(self):
        summary = self.get("risk-list/")["summary"]
        tile = self.tiles()["risk"]
        self.assertEqual((tile["value"], tile["urgent"]), (summary["total"], summary["urgent"]))
        self.assertGreater(tile["value"], 0)  # должник из setUp

    def test_forecast_is_the_forecast_report(self):
        forecast = self.client.get(f"{API}forecast/").data["forecast"]
        tile = self.tiles()["forecast"]
        self.assertEqual(tile["status"], forecast["status"])
        self.assertEqual(tile["status"], "range")
        self.assertEqual((tile["low"], tile["high"]), (str(forecast["low"]), str(forecast["high"])))

    def test_one_broken_tile_does_not_break_the_screen(self):
        with (
            mock.patch("domains.platform.analytics.dashboard.risk_list", side_effect=RuntimeError),
            self.assertLogs("domains.platform.analytics.dashboard", "ERROR"),
        ):
            tiles = self.tiles()
        self.assertIsNone(tiles["risk"])
        self.assertEqual(tiles["revenue"]["value"], "45000")


class RolePermissionTests(DashboardFixtures):
    """ТЗ п. 11.7 и п. 2: права ролей на всех отчётах M3."""

    def assert_status(self, user, status, **params):
        self.client.force_authenticate(user)
        for path, extra in ANALYTICS_ENDPOINTS:
            with self.subTest(role=user.role, path=path):
                response = self.client.get(f"{API}{path}", {**extra, **params})
                self.assertEqual(response.status_code, status, response.content[:200])

    def test_admin_teacher_accountant_do_not_see_org_summary(self):
        roles = (User.Role.ADMIN, User.Role.TEACHER, User.Role.ACCOUNTANT)
        for index, role in enumerate(roles):
            user = self.user(f"+7703000000{index}", role, branches=[self.abaya])
            self.assert_status(user, 403)

    def test_manager_sees_only_own_branches(self):
        manager = self.user("+77030000011", User.Role.MANAGER, branches=[self.abaya])
        self.assert_status(manager, 200)
        tiles = self.tiles()
        self.assertEqual(tiles["revenue"]["value"], "30000")  # без Саина
        self.assertEqual(tiles["group_fill"]["capacity"], 20)
        data = self.get("dashboard/")
        self.assertFalse(data["all_branches"])
        self.assertEqual([b["name"] for b in data["branches"]], ["Абая"])
        # Чужой филиал не открыть ни в одном отчёте, где выбирается филиал.
        self.client.force_authenticate(manager)
        for path, extra in ANALYTICS_ENDPOINTS:
            if path == "catalog/":
                continue
            with self.subTest(path=path):
                response = self.client.get(f"{API}{path}", {**extra, "branch": str(self.saina.pk)})
                self.assertEqual(response.status_code, 403, response.content[:200])

    def test_admin_with_org_summary_setting_sees_own_branch_only(self):
        self.org.settings = {**(self.org.settings or {}), ADMIN_SEES_ORG_SUMMARY: True}
        self.org.save(update_fields=["settings"])
        admin = self.user("+77030000012", User.Role.ADMIN, branches=[self.saina])
        self.client.force_authenticate(admin)
        self.assertEqual(self.tiles()["revenue"]["value"], "15000")


class ExportEveryReportTests(DashboardFixtures):
    def test_every_m3_report_exports_to_excel(self):
        self.assertGreaterEqual(len(REPORTS), 13)  # обзор + двенадцать отчётов
        for name in REPORTS:
            with self.subTest(report=name):
                response = self.client.get(f"{API}export/", {**self.period(), "report": name})
                self.assertEqual(response.status_code, 200, response.content[:200])
                self.assertIn("spreadsheetml", response["Content-Type"])
                book = openpyxl.load_workbook(io.BytesIO(response.content))
                self.assertTrue(book.sheetnames)

    def test_overview_export_starts_with_dashboard_tiles(self):
        response = self.client.get(f"{API}export/", {**self.period(), "report": "overview"})
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames[0], "Дашборд")
        values = [row[0].value for row in book["Дашборд"].iter_rows() if row and row[0].value]
        for label in ("Выручка", "Задолженность", "Конверсия заявок", "Прогноз выручки"):
            self.assertIn(label, values)


@tag("tenant_isolation")
class DashboardTenantIsolationTests(DashboardFixtures):
    def test_other_center_is_not_counted(self):
        other = Organization.objects.create(name="Другой", slug="dash-other")
        branch = Branch.objects.create(organization=other, name="Чужой")
        direction = Direction.objects.create(organization=other, name="Танцы")
        Group.objects.create(
            organization=other, branch=branch, direction=direction, name="Чужая", capacity=50
        )
        tiles = self.tiles()
        self.assertEqual(tiles["group_fill"]["capacity"], 24)
        self.assertEqual(tiles["revenue"]["value"], "45000")
        stranger = User.objects.create_user(
            phone="+77030000099",
            password="x",
            full_name="Чужой владелец",
            organization=other,
            role=User.Role.OWNER,
        )
        self.client.force_authenticate(stranger)
        tiles = self.tiles()
        self.assertEqual(tiles["revenue"]["value"], 0)
        self.assertEqual(tiles["group_fill"]["capacity"], 50)
        self.assertEqual(tiles["lead_conversion"]["leads"], 0)
