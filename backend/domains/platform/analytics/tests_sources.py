"""Отчёт по источникам заявок (TRU-116)."""

import io

import openpyxl

from domains.money.subscriptions.models import Subscription
from domains.money.subscriptions.subscription_types import create_type
from domains.people.clients.models import Child
from domains.platform.leads.models import Lead, LeadSource
from domains.platform.tenants.models import Direction

from .tests_funnel import FunnelFixtures, S

URL = "/api/v1/analytics/sources/"


class SourcesTests(FunnelFixtures):
    def setUp(self):
        super().setUp()
        self.word = LeadSource.objects.create(organization=self.org, name="Рекомендации")
        self.type = create_type(
            self.org, name="8 занятий", price=30000, quota_sessions=8, duration_days=30
        )

    def sold(self, lead, price):
        child = Child.objects.create(
            organization=self.org, full_name="Ребёнок", birth_date="2018-01-01", gender="female"
        )
        subscription = Subscription.objects.create(
            organization=self.org,
            child=child,
            subscription_type_version=self.type.versions.latest(),
            direction=self.ballet,
            branch=self.abaya,
            starts_on=self.today,
            ends_on=self.today,
            list_price=price,
            price=price,
        )
        Lead.objects.filter(pk=lead.pk).update(sold_subscription=subscription)

    def test_quality_not_just_quantity(self):
        for _ in range(3):
            self.lead(S.NEW, source=self.instagram)
        bought = self.lead(S.NEW, S.CONTACTED, S.PURCHASED, source=self.instagram)
        self.sold(bought, 30000)
        for price in (40000, 20000):
            lead = self.lead(
                S.NEW,
                S.CONTACTED,
                S.TRIAL_SCHEDULED,
                S.TRIAL_ATTENDED,
                S.PURCHASED,
                source=self.word,
            )
            self.sold(lead, price)
        items = {i["label"]: i for i in self.client.get(URL).data["items"]}
        insta, word = items["Instagram"], items["Рекомендации"]
        self.assertEqual((insta["leads"], insta["purchased"], insta["conversion"]), (4, 1, 25.0))
        self.assertEqual((word["leads"], word["purchased"], word["conversion"]), (2, 2, 100.0))
        self.assertEqual((word["trial"], word["trial_rate"]), (2, 100.0))
        self.assertEqual((word["avg_check"], insta["avg_check"]), ("30000", "30000"))
        # Больше покупок — выше, хоть заявок и меньше.
        self.assertEqual(self.client.get(URL).data["items"][0]["label"], "Рекомендации")
        self.assertTrue(word["small_sample"])

    def test_renewals_excluded_and_direction_filter(self):
        dance = Direction.objects.create(organization=self.org, name="Танцы")
        self.lead(S.NEW, source=self.instagram, kind="renewal")
        other = self.lead(S.NEW, source=self.instagram)
        self.lead(S.NEW, source=self.instagram)
        Lead.objects.filter(pk=other.pk).update(direction=dance)
        self.assertEqual(self.client.get(URL).data["items"][0]["leads"], 2)
        filtered = self.client.get(URL, {"direction": str(dance.pk)}).data["items"]
        self.assertEqual([(i["label"], i["leads"]) for i in filtered], [("Instagram", 1)])

    def test_by_month_and_export(self):
        self.lead(S.NEW, S.CONTACTED, S.PURCHASED, source=self.instagram)
        data = self.client.get(URL).data
        self.assertEqual(
            [(r["label"], r["leads"], r["purchased"]) for r in data["by_month"]],
            [("Instagram", 1, 1)],
        )
        response = self.client.get("/api/v1/analytics/export/", {"report": "sources"})
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Источники", "По месяцам"])
        values = [c.value for row in book["Источники"].iter_rows() for c in row]
        self.assertIn("Instagram", values)
