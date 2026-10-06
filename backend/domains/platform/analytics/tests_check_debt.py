"""Средний чек и динамика задолженности (TRU-124)."""

import io
from datetime import datetime, time, timedelta
from decimal import Decimal

import openpyxl
from django.test import tag

from domains.money.subscriptions.debt import debt_structure, debt_total, repaid_debt
from domains.money.subscriptions.models import Subscription
from domains.platform.leads.models import Lead, LeadKind, LeadSource
from domains.platform.tenants.models import Organization
from domains.platform.users.models import User

from .check_debt import average_check, debt_dynamics, distribution
from .models import MetricSnapshot
from .period import Period, period_for
from .scope import Scope
from .tests import AnalyticsFixtures

CHECK_URL = "/api/v1/analytics/average-check/"
DEBT_URL = "/api/v1/analytics/debt/"


class CheckFixtures(AnalyticsFixtures):
    def sale(self, branch, price, *, day=None, discount=0, reason="", child=None, **extra):
        subscription = self.subscription(branch, child=child, price=price)
        moment = self.tz.localize(datetime.combine(day or self.today, time(12)))
        Subscription.objects.filter(pk=subscription.pk).update(
            created_at=moment,
            list_price=price + discount,
            discount_amount=discount,
            discount_reason=reason,
            **extra,
        )
        subscription.refresh_from_db()
        return subscription

    def aged(self, subscription, days):
        """Абонемент начался `days` дней назад — давность долга как на экране."""
        Subscription.objects.filter(pk=subscription.pk).update(
            starts_on=self.today - timedelta(days=days)
        )


class AverageCheckTests(CheckFixtures):
    def test_average_with_median_and_one_expensive_outlier(self):
        """Один годовой абонемент тянет среднее вверх, медиана остаётся на месте."""
        for price in (20000, 30000, 30000):
            self.sale(self.abaya, price)
        self.sale(self.saina, 240000)
        data = average_check(Scope(self.org, None), period_for("month", self.today))
        summary = data["summary"]
        self.assertEqual((summary["count"], summary["amount"]), (4, "320000"))
        self.assertEqual((summary["average"], summary["median"]), ("80000", "30000"))
        self.assertEqual((summary["min"], summary["max"]), ("20000", "240000"))
        # Квартили с интерполяцией: половина абонементов — от 27 500 до 82 500.
        self.assertEqual((summary["p25"], summary["p75"]), ("27500", "82500"))
        # Плитка «Средний чек» — та же цифра.
        self.assertEqual(self.metric("average_check")["value"], summary["average"])
        self.assertEqual(sum(b["count"] for b in data["distribution"]), 4)

    def test_distribution_bins(self):
        bins = distribution({Decimal(20000): 1, Decimal(30000): 2, Decimal(240000): 1})
        self.assertEqual(bins[0]["from"], "0")
        self.assertEqual((bins[0]["to"], bins[0]["count"]), ("50000", 3))
        self.assertEqual(bins[-1]["count"], 1)
        self.assertLessEqual(len(bins), 8)
        self.assertEqual(distribution({Decimal(30000): 3})[0]["count"], 3)
        self.assertEqual(distribution({}), [])

    def test_breakdowns_by_branch_client_and_source(self):
        instagram = LeadSource.objects.create(organization=self.org, name="Instagram")
        first = self.sale(self.abaya, 30000)
        Lead.objects.create(
            organization=self.org,
            parent_name="Мама",
            phone="+77011111111",
            kind=LeadKind.NEW,
            source=instagram,
            converted_child=first.child,
        )
        self.sale(self.abaya, 40000, child=first.child, renewed_from=first)
        self.sale(self.saina, 20000)
        by = average_check(Scope(self.org, None), period_for("month", self.today))["by"]
        branches = {item["label"]: (item["count"], item["average"]) for item in by["branch"]}
        self.assertEqual(branches, {"Абая": (2, "35000"), "Саина": (1, "20000")})
        clients = {item["key"]: item["average"] for item in by["client"]}
        self.assertEqual(clients, {"new": "25000", "renewal": "40000"})
        sources = {item["label"]: item["count"] for item in by["source"]}
        self.assertEqual(sources, {"Instagram": 2, None: 1})

    def test_discount_effect_and_reasons(self):
        self.sale(self.abaya, 25000, discount=5000, reason="second_child")
        self.sale(self.abaya, 27000, discount=3000, reason="promotion")
        self.sale(self.abaya, 30000)
        discounts = average_check(Scope(self.org, None), period_for("month", self.today))[
            "discounts"
        ]
        self.assertEqual((discounts["total"], discounts["count"]), ("8000", 2))
        self.assertEqual((discounts["list_average"], discounts["average"]), ("30000", "27333.3"))
        self.assertEqual(discounts["effect"], "2666.7")
        self.assertEqual(
            [(r["key"], r["label"], r["amount"]) for r in discounts["reasons"]],
            [("second_child", "Второй ребёнок", "5000"), ("promotion", "Акция", "3000")],
        )
        self.assertEqual(self.metric("discount_total")["value"], "8000")

    def test_sale_date_and_monthly_dynamics(self):
        month = period_for("month", self.today)
        last_month_day = month.start - timedelta(days=1)
        self.sale(self.abaya, 30000)
        self.sale(self.abaya, 50000, day=last_month_day)
        data = average_check(Scope(self.org, None), month)
        self.assertEqual(data["summary"]["count"], 1)
        monthly = data["monthly"]
        self.assertEqual(len(monthly), 12)
        self.assertEqual(monthly[-1]["month"], month.start.isoformat())
        self.assertEqual((monthly[-1]["count"], monthly[-1]["average"]), (1, "30000"))
        self.assertEqual(monthly[-2]["average"], "50000")

    def test_api_scope_and_permissions(self):
        self.sale(self.abaya, 30000)
        self.sale(self.saina, 10000)
        manager = self.user("+77010000002", User.Role.MANAGER, [self.abaya])
        self.client.force_authenticate(manager)
        data = self.client.get(CHECK_URL).data
        self.assertEqual(data["summary"]["count"], 1)
        self.assertEqual(data["period"]["preset"], "month")
        self.assertEqual(
            self.client.get(CHECK_URL, {"branch": str(self.saina.pk)}).status_code, 403
        )
        teacher = self.user("+77010000005", User.Role.TEACHER)
        self.client.force_authenticate(teacher)
        self.assertEqual(self.client.get(CHECK_URL).status_code, 403)
        self.assertEqual(self.client.get(DEBT_URL).status_code, 403)

    @tag("tenant_isolation")
    def test_other_organization_is_invisible(self):
        self.sale(self.abaya, 30000)
        other = Organization.objects.create(name="Чужой", slug="other-check")
        stranger = User.objects.create_user(
            phone="+77019999998", password="x", full_name="Ч", organization=other, role="owner"
        )
        self.client.force_authenticate(stranger)
        self.assertEqual(self.client.get(CHECK_URL).data["summary"]["count"], 0)
        self.assertEqual(self.client.get(DEBT_URL).data["end"]["total"], "0")


class DebtDynamicsTests(CheckFixtures):
    def test_total_is_the_debtors_screen_total(self):
        self.pay(self.subscription(self.abaya), 10000)
        overpaid = self.subscription(self.saina)
        self.pay(overpaid, 35000)
        self.subscription(self.saina, price=12000)
        self.client.force_authenticate(self.owner)
        screen = self.client.get("/api/v1/subscriptions/debtors/").data["total_debt"]
        end = self.client.get(DEBT_URL).data["end"]
        self.assertEqual((end["total"], end["source"]), ("32000", "live"))
        self.assertEqual(Decimal(screen), Decimal(end["total"]))
        saina_screen = self.client.get(
            "/api/v1/subscriptions/debtors/", {"branch": str(self.saina.pk)}
        ).data["total_debt"]
        saina = self.client.get(DEBT_URL, {"branch": str(self.saina.pk)}).data["end"]
        self.assertEqual(Decimal(saina_screen), Decimal(saina["total"]))

    def test_age_structure_and_debtors_share(self):
        fresh, month_old, old = (self.subscription(self.abaya) for _ in range(3))
        self.aged(fresh, 10)
        self.aged(month_old, 45)
        self.aged(old, 90)
        self.pay(old, 25000)
        self.pay(self.subscription(self.abaya), 30000)  # оплачен — не должник
        structure = debt_structure(self.org)
        self.assertEqual(
            structure["by_age"],
            {"0_30": Decimal(30000), "31_60": Decimal(30000), "over_60": Decimal(5000)},
        )
        self.assertEqual(structure["total"], debt_total(self.org))
        end = debt_dynamics(Scope(self.org, None), period_for("month", self.today))["end"]
        self.assertEqual(end["by_age"], {"0_30": "30000", "31_60": "30000", "over_60": "5000"})
        # 3 должника из 4 активных детей.
        self.assertEqual((end["debtors_active"], end["active_children"]), (3, 4))
        self.assertEqual(end["debtors_share"], "75")
        self.assertEqual(self.metric("debtors_share")["value"], "75")

    def test_past_period_uses_snapshots_not_recount(self):
        self.subscription(self.abaya)  # сегодняшний долг 30000
        month = period_for("month", self.today)
        past = Period(month.start - timedelta(days=60), month.start - timedelta(days=31))
        for day, value in ((past.end, 50000), (past.previous().end, 40000)):
            MetricSnapshot.objects.create(
                organization=self.org, metric="debt_total", date=day, value=value
            )
        MetricSnapshot.objects.create(
            organization=self.org, metric="debt_age_over_60", date=past.end, value=20000
        )
        data = debt_dynamics(Scope(self.org, None), past)
        self.assertEqual((data["end"]["source"], data["end"]["total"]), ("snapshot", "50000"))
        self.assertEqual(data["end"]["by_age"]["over_60"], "20000")
        self.assertEqual((data["previous"]["total"], data["change"]), ("40000", "10000"))
        self.assertEqual(data["change_percent"], "25")
        self.assertEqual(data["history_since"], past.previous().end.isoformat())

    def test_no_snapshot_means_no_number(self):
        self.subscription(self.abaya)
        month = period_for("month", self.today)
        data = debt_dynamics(Scope(self.org, None), month)
        self.assertEqual(data["end"]["total"], "30000")
        self.assertIsNone(data["previous"]["total"])
        self.assertIsNone(data["change"])
        self.assertEqual(len(data["monthly"]), 12)
        self.assertEqual(data["monthly"][-1]["total"], "30000")
        self.assertIsNone(data["monthly"][0]["total"])

    def test_snapshots_of_two_branches_add_up(self):
        month = period_for("month", self.today)
        day = month.start - timedelta(days=40)
        for branch, value in ((self.abaya, 10000), (self.saina, 5000)):
            MetricSnapshot.objects.create(
                organization=self.org, metric="debt_total", branch=branch, date=day, value=value
            )
        scope = Scope(self.org, (self.abaya.pk, self.saina.pk))
        point = debt_dynamics(scope, Period(day, day))["end"]
        self.assertEqual(point["total"], "15000")
        self.assertIsNone(point["debtors_share"])

    def test_repaid_part_of_old_debt(self):
        month = period_for("month", self.today)
        before = month.start - timedelta(days=10)
        old = self.sale(self.abaya, 30000, day=before)
        self.pay(old, 10000, day=before)  # на начало периода долг 20000
        self.pay(old, 5000)  # погасили 5000
        overpaid = self.sale(self.abaya, 10000, day=before)
        self.pay(overpaid, 25000)  # переплата: погашено только 10000
        self.pay(self.sale(self.abaya, 30000), 30000)  # новая продажа — не старый долг
        start, end = month.bounds(self.org)
        self.assertEqual(repaid_debt(self.org, start, end), (Decimal(30000), Decimal(15000)))
        repaid = debt_dynamics(Scope(self.org, None), month)["repaid"]
        self.assertEqual(
            (repaid["owed_at_start"], repaid["repaid"], repaid["remaining"], repaid["percent"]),
            ("30000", "15000", "15000", "50"),
        )


class CheckDebtExportTests(CheckFixtures):
    def test_export_has_check_and_debt_sheets(self):
        self.sale(self.abaya, 30000, discount=5000, reason="promotion")
        self.subscription(self.saina, price=12000)
        self.client.force_authenticate(self.owner)
        response = self.client.get(
            "/api/v1/analytics/export/", {"report": "check_debt", "period": "month"}
        )
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        for name in ("Средний чек", "Распределение цен", "Скидки по причинам", "Задолженность"):
            self.assertIn(name, book.sheetnames)
        sheet = book["Задолженность"]
        values = [row for row in sheet.iter_rows(values_only=True) if row and row[0]]
        end_row = next(row for row in values if str(row[0]).startswith("Конец периода"))
        self.assertEqual(end_row[1], 42000)
        reasons = [row for row in book["Скидки по причинам"].iter_rows(values_only=True)]
        self.assertIn(("Акция", 1, 5000), reasons)
