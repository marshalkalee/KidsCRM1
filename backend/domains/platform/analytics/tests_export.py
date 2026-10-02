"""Выгрузка любого отчёта в Excel (TRU-114)."""

import io
from datetime import date

import openpyxl

from .tests import AnalyticsFixtures

URL = "/api/v1/analytics/export/"


class ExportTests(AnalyticsFixtures):
    def setUp(self):
        super().setUp()
        self.pay(self.subscription(self.abaya), 30000)
        self.pay(self.subscription(self.saina), 15000)
        self.client.force_authenticate(self.owner)

    def book(self, **params):
        response = self.client.get(URL, params)
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.assertIn("spreadsheetml", response["Content-Type"])
        return openpyxl.load_workbook(io.BytesIO(response.content))

    def table(self, sheet, first_column):
        """Заголовок таблицы и строки под ним (шапка отчёта — выше)."""
        rows = list(sheet.iter_rows())
        head = next(i for i, row in enumerate(rows) if row[0].value == first_column)
        return rows[head], rows[head + 1 :]

    def test_every_report_exports(self):
        for report in ("overview", "revenue", "attendance", "funnel"):
            with self.subTest(report=report):
                book = self.book(report=report, period="month")
                self.assertGreater(len(book.sheetnames), 2)

    def test_header_says_what_it_is(self):
        sheet = self.book(report="revenue", branch=str(self.abaya.pk))["Показатели"]
        header = [sheet.cell(row=i, column=1).value for i in range(1, 7)]
        self.assertEqual(header[0], "Выручка по оплатам")
        self.assertIn("Центр: True Ballet", header)
        self.assertIn("Филиалы: Абая", header)
        self.assertTrue(any(str(line).startswith("Период: ") for line in header))
        self.assertTrue(any(str(line).startswith("Выгружено: ") for line in header))

    def test_same_numbers_as_screen_and_real_numbers(self):
        screen = self.client.get(
            "/api/v1/analytics/breakdown/", {"metric": "revenue", "by": "branch"}
        ).data["items"]
        book = self.book(report="revenue")
        head, rows = self.table(book["По филиалам"], "Филиал")
        self.assertEqual(head[1].value, "Выручка, ₸")
        values = {row[0].value: row[1].value for row in rows if row[0].value != "Итого"}
        self.assertEqual(values, {item["label"]: int(item["value"]) for item in screen})
        self.assertIsInstance(values["Абая"], int)
        # Итог — формулой, а не текстом: Excel пересчитает сам.
        total = [row for row in rows if row[0].value == "Итого"][0]
        first, last = head[1].row + 1, total[1].row - 1
        self.assertEqual(total[1].value, f"=SUM(B{first}:B{last})")  # ровно строки данных
        self.assertIsNone(total[2].value)  # доли не складываем
        self.assertTrue(head[0].font.bold)
        self.assertEqual(head[0].fill.fgColor.rgb[-6:], "FFE9EC")

    def test_dates_are_dates(self):
        book = self.book(report="revenue", period="month")
        _, rows = self.table(book["Динамика"], "Период с")
        self.assertIsInstance(rows[0][0].value.date(), date)
        self.assertEqual(rows[0][0].number_format, "DD.MM.YYYY")

    def test_unknown_report_and_permissions(self):
        self.assertEqual(self.client.get(URL, {"report": "nope"}).status_code, 400)
        admin = self.user("+77010000099", "admin")
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.get(URL, {"report": "revenue"}).status_code, 403)
