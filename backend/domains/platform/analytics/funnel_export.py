"""Выгрузка воронки в Excel (TRU-115): этапы и разрезы, как на экране."""

import io

import openpyxl
from openpyxl.styles import Font

from .funnel import STAGES

BY_TITLES = {
    "source": "Источник",
    "direction": "Направление",
    "branch": "Филиал",
    "manager": "Ответственный",
}


def _header(sheet, titles, widths):
    sheet.append(titles)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def funnel_workbook(period, summary, breakdowns) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Воронка"
    _header(
        sheet,
        ["Этап", "Заявок", "% от заявок", "% от прошлого этапа", "Сейчас на этапе"],
        [26, 12, 14, 20, 18],
    )
    total = summary["total"]
    previous = None
    for stage in summary["stages"]:
        sheet.append(
            [
                stage["label"],
                stage["count"],
                round(stage["count"] * 100 / total, 1) if total else None,
                round(stage["count"] * 100 / previous, 1) if previous else None,
                stage["current"],
            ]
        )
        previous = stage["count"]
    sheet.append([])
    sheet.append(["Купили без пробного", summary["purchased_without_trial"]])
    sheet.append(["Думают", summary["thinking"]])
    sheet.append(["Отказались", summary["rejected"]])
    sheet.append(["Период", f"{period.start:%d.%m.%Y} – {period.end:%d.%m.%Y}"])
    sheet.append(["Продления", "не входят: отдельная воронка"])

    for dimension, items in breakdowns.items():
        part = workbook.create_sheet(BY_TITLES[dimension])
        _header(
            part,
            [BY_TITLES[dimension], *[stage.label for stage in STAGES], "Конверсия, %"],
            [26, 12, 12, 14, 14, 12, 14],
        )
        for item in items:
            part.append(
                [
                    item["label"] or "Не указано",
                    *[item["stages"][stage] for stage in STAGES],
                    item["conversion"],
                ]
            )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
