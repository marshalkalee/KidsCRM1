"""Конверсия продлений (TRU-126, ТЗ п. 5.3)."""

import io
from datetime import date, time, timedelta

import openpyxl
from django.test import tag

from domains.money.subscriptions.models import Subscription
from domains.money.subscriptions.renewal_conversion import add_months
from domains.money.subscriptions.subscription_types import create_type
from domains.people.clients.models import Child
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.users.models import User
from domains.scheduling.groups.models import Group, GroupMembership
from domains.scheduling.schedule_templates.models import ScheduleTemplate, ScheduleTemplateSlot

from .period import Period
from .renewal_report import renewal_conversion_report
from .scope import Scope
from .tests_forecast import ForecastFixtures

URL = "/api/v1/analytics/renewal-conversion/"
EXPORT_URL = "/api/v1/analytics/export/"


class RenewalFixtures(ForecastFixtures):
    def setUp(self):
        super().setUp()
        # Месяц, у которого окно продления уже точно прошло.
        self.past = add_months(self.month, -2)
        self.past_end = add_months(self.past, 1) - timedelta(days=1)
        self.past_period = Period(self.past, self.past_end, "custom")

    def ended_sub(self, ends_on, **kwargs):
        kwargs.setdefault("status", Subscription.Status.EXPIRED)
        return self.sub(starts_on=ends_on - timedelta(days=30), ends_on=ends_on, **kwargs)

    def renew(self, old, *, after_days=3, bought_after=None, **kwargs):
        start = old.ends_on + timedelta(days=after_days)
        bought = old.ends_on + timedelta(days=after_days if bought_after is None else bought_after)
        return self.sub(
            child=old.child,
            starts_on=start,
            ends_on=start + timedelta(days=30),
            direction=kwargs.pop("direction", old.direction),
            created_on=min(bought, self.today),
            **kwargs,
        )

    def report(self, period=None, branch_ids=None):
        return renewal_conversion_report(Scope(self.org, branch_ids), period or self.past_period)


class ManualCountTests(RenewalFixtures):
    def test_month_matches_manual_count(self):
        """Ручной подсчёт за месяц: закончилось 5, продлили 2 — 40%."""
        day = self.past + timedelta(days=9)
        self.renew(self.ended_sub(day))  # продлил через 3 дня
        button = self.ended_sub(day)
        # Кнопкой «Продлить» — продление, даже после долгой паузы.
        self.renew(button, after_days=40, renewed_from=button)
        self.renew(self.ended_sub(day), direction=self.stretch)  # другое направление
        self.renew(self.ended_sub(day), after_days=20)  # позже окна
        self.ended_sub(day)  # не вернулся
        self.ended_sub(self.past - timedelta(days=1))  # закончился до периода
        self.ended_sub(day, branch=self.saina)  # другой филиал — только в общей цифре

        summary = self.report(branch_ids=(self.abaya.pk,))["summary"]

        self.assertEqual(summary["ended"], 5)
        self.assertEqual(summary["decided"], 5)
        self.assertEqual(summary["renewed"], 2)
        self.assertEqual(summary["lost"], 3)
        self.assertEqual(summary["rate"], 40.0)
        self.assertEqual(summary["pending"], 0)
        self.assertEqual(self.report()["summary"]["ended"], 6)

    def test_open_window_is_not_counted_as_lost(self):
        """Закончился три дня назад — в конверсию не идёт, даже если уже продлил:
        иначе ранние продления без поздних завышали бы процент."""
        ends = self.today - timedelta(days=3)
        self.ended_sub(ends)
        self.renew(self.ended_sub(ends), after_days=1)
        self.ended_sub(self.today - timedelta(days=40))  # окно прошло, не продлил

        period = Period(self.today - timedelta(days=60), self.today, "custom")
        summary = self.report(period)["summary"]

        self.assertEqual(summary["ended"], 3)
        self.assertEqual(summary["pending"], 2)
        self.assertEqual(summary["pending_renewed"], 1)
        self.assertEqual(summary["decided"], 1)
        self.assertEqual(summary["renewed"], 0)
        self.assertEqual(summary["rate"], 0.0)

    def test_first_renewals_are_separate(self):
        day = self.past + timedelta(days=5)
        # Ребёнок 1: первый абонемент продлил, продление тоже закончилось и продлено.
        first = self.ended_sub(day - timedelta(days=33))
        second = self.renew(first)  # заканчивается в периоде — последующее
        self.renew(second)
        # Ребёнок 2: первый абонемент закончился в периоде, не продлил.
        self.ended_sub(day)
        # Ребёнок 3: первый в периоде, продлил.
        self.renew(self.ended_sub(day))

        summary = self.report()["summary"]

        self.assertEqual(summary["first"], {"decided": 2, "renewed": 1, "rate": 50.0})
        self.assertEqual(summary["repeat"], {"decided": 1, "renewed": 1, "rate": 100.0})

    def test_grace_days_come_from_org_setting_and_reach_forecast(self):
        old = self.ended_sub(self.past + timedelta(days=5))
        self.renew(old, after_days=20)
        self.assertEqual(self.report()["summary"]["renewed"], 0)

        self.org.settings = {**self.org.settings, "renewal_grace_days": 30}
        self.org.save()

        data = self.report()
        self.assertEqual(data["summary"]["renewed"], 1)
        self.assertEqual(data["rules"]["grace_days"], 30)
        self.assertEqual(self.forecast()["rules"]["grace_days"], 30)


class BreakdownTests(RenewalFixtures):
    def test_renewal_delay_buckets(self):
        self.org.settings = {**self.org.settings, "renewal_grace_days": 30}
        self.org.save()
        day = self.past + timedelta(days=5)
        self.renew(self.ended_sub(day), after_days=1, bought_after=-2)  # купили заранее
        self.renew(self.ended_sub(day), after_days=5)
        self.renew(self.ended_sub(day), after_days=20)

        gaps = self.report()["gaps"]

        values = {row["key"]: row["value"] for row in gaps["buckets"]}
        self.assertEqual(values, {"early": 1, "week": 1, "two_weeks": 0, "month": 1, "later": 0})
        self.assertEqual(gaps["average_days"], 8.3)  # (0 + 5 + 20) / 3
        self.assertEqual(gaps["long_pause"], 1)

    def test_type_and_age(self):
        day = self.past + timedelta(days=5)
        unlimited = create_type(
            self.org, name="Безлимит", price=45000, is_unlimited=True, duration_days=30
        )
        young = Child.objects.create(
            organization=self.org, full_name="Малыш", birth_date=date(2021, 1, 1), gender="male"
        )
        self.renew(self.ended_sub(day, child=young, type_=unlimited))
        self.ended_sub(day)  # ребёнок 2018 г. р.

        data = self.report()["breakdowns"]

        types = {row["label"]: (row["ended"], row["renewed"]) for row in data["type"]}
        self.assertEqual(types, {"Безлимит": (1, 1), "8 занятий": (1, 0)})
        ages = [row["key"] for row in data["age"]]
        self.assertEqual(ages, sorted(ages, key=int))
        self.assertEqual(len(ages), 2)

    def test_teacher_breakdown_has_context_and_is_not_a_ranking(self):
        teacher_b = self.user("+77010000011", User.Role.TEACHER)
        teacher_b.full_name = "Бота"
        teacher_b.save()
        teacher_a = self.user("+77010000012", User.Role.TEACHER)
        teacher_a.full_name = "Айгуль"
        teacher_a.save()
        morning = self.group("Утро", teacher_a, time(10), capacity=10)
        evening = self.group("Вечер", teacher_b, time(18), capacity=8)
        day = self.past + timedelta(days=5)
        for group, renewed in ((morning, False), (morning, False), (evening, True)):
            old = self.ended_sub(day)
            self.member(group, old.child, joined=old.starts_on - timedelta(days=60))
            if renewed:
                self.renew(old)
        self.ended_sub(day)  # без группы

        data = self.report()["breakdowns"]

        teachers = data["teacher"]
        # По алфавиту, а не по конверсии; без преподавателя — в конце.
        self.assertEqual(
            [row["label"] for row in teachers], ["Айгуль", "Бота", "Без преподавателя"]
        )
        self.assertEqual(teachers[0]["rate"], 0.0)
        self.assertTrue(teachers[0]["small"])
        context = teachers[0]["context"]
        self.assertEqual(context["fill_percent"], 20)  # 2 из 10
        self.assertEqual(context["groups"][0]["schedule"], "Пн 10:00")
        groups = {row["label"]: row["ended"] for row in data["group"]}
        self.assertEqual(groups, {"Утро": 2, "Вечер": 1, "Без группы": 1})

    def group(self, name, teacher, starts, capacity):
        group = Group.objects.create(
            organization=self.org,
            branch=self.abaya,
            direction=self.ballet,
            name=name,
            capacity=capacity,
        )
        group.teachers.add(teacher)
        template = ScheduleTemplate.objects.create(
            organization=self.org, group=group, valid_from=self.today - timedelta(days=400)
        )
        ScheduleTemplateSlot.objects.create(
            organization=self.org, template=template, weekday=0, start_time=starts, teacher=teacher
        )
        return group

    def member(self, group, child, joined):
        GroupMembership.objects.create(
            organization=self.org, group=group, child=child, joined_at=joined
        )

    def test_trend_has_twelve_months_and_marks_unsettled(self):
        self.renew(self.ended_sub(self.past + timedelta(days=3)))

        trend = self.report(Period(self.month, self.today, "month"))["trend"]

        self.assertEqual(len(trend), 12)
        self.assertEqual(trend[-1]["month"], self.month.isoformat())
        self.assertFalse(trend[-1]["complete"])
        past = next(row for row in trend if row["month"] == self.past.isoformat())
        self.assertEqual((past["ended"], past["renewed"], past["complete"]), (1, 1, True))


class ApiTests(RenewalFixtures):
    def test_api_and_export(self):
        self.renew(self.ended_sub(self.past + timedelta(days=3)))
        self.client.force_authenticate(self.owner)
        params = {
            "period": "custom",
            "from": self.past.isoformat(),
            "to": self.past_end.isoformat(),
        }

        response = self.client.get(URL, params)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.data["summary"]["renewed"], 1)

        response = self.client.get(EXPORT_URL, {**params, "report": "renewal_conversion"})
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames[:3], ["Итог", "По месяцам", "Срок продления"])
        self.assertIn("По преподавателям", book.sheetnames)

    def test_teacher_has_no_access(self):
        teacher = self.user("+77010000021", User.Role.TEACHER)
        self.client.force_authenticate(teacher)
        self.assertEqual(self.client.get(URL).status_code, 403)

    def test_settings_api_exposes_grace_days(self):
        self.client.force_authenticate(self.owner)
        data = self.client.get("/api/v1/organization/settings/").json()
        self.assertEqual(data["renewal_grace_days"], 14)


@tag("tenant_isolation")
class RenewalTenantIsolationTests(RenewalFixtures):
    def test_other_center_is_not_counted(self):
        other = Organization.objects.create(name="Другой", slug="renewal-other")
        branch = Branch.objects.create(organization=other, name="Чужой")
        direction = Direction.objects.create(organization=other, name="Балет")
        other_type = create_type(other, name="8", price=50000, quota_sessions=8, duration_days=30)
        child = Child.objects.create(
            organization=other, full_name="Чужой", birth_date=date(2018, 1, 1), gender="female"
        )
        ends = self.past + timedelta(days=5)
        Subscription.objects.create(
            organization=other,
            child=child,
            subscription_type_version=other_type.versions.latest(),
            direction=direction,
            branch=branch,
            starts_on=ends - timedelta(days=30),
            ends_on=ends,
            list_price=50000,
            price=50000,
        )
        self.ended_sub(ends)
        self.client.force_authenticate(self.owner)
        params = {
            "period": "custom",
            "from": self.past.isoformat(),
            "to": self.past_end.isoformat(),
        }

        data = self.client.get(URL, params).data
        self.assertEqual(data["summary"]["ended"], 1)

        # Чужой филиал в запросе — не данные чужого центра, а отказ.
        response = self.client.get(URL, {**params, "branch": str(branch.pk)})
        self.assertEqual(response.status_code, 403)
