"""Прогноз выручки от активных абонементов (TRU-125)."""

import io
from datetime import datetime, time, timedelta
from decimal import Decimal

import openpyxl
from django.test import tag

from domains.money.subscriptions.models import Subscription
from domains.money.subscriptions.renewal_conversion import (
    RENEWAL_GRACE_DAYS,
    RenewalIndex,
    add_months,
    load_rows,
    renewal_conversion,
)
from domains.money.subscriptions.subscription_types import create_type
from domains.people.clients.models import Child
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.users.models import User

from .forecast import MIN_SAMPLE_POINT, MIN_SAMPLE_RANGE, calibrate, revenue_forecast, wilson
from .scope import Scope
from .tests import AnalyticsFixtures

URL = "/api/v1/analytics/forecast/"


class ForecastFixtures(AnalyticsFixtures):
    def setUp(self):
        super().setUp()
        self.month = self.today.replace(day=1)
        self.next_month = add_months(self.month, 1)
        self.stretch = Direction.objects.create(organization=self.org, name="Растяжка")

    def sub(
        self,
        *,
        starts_on,
        ends_on=None,
        child=None,
        price=30000,
        created_on=None,
        direction=None,
        branch=None,
        renewed_from=None,
        status=Subscription.Status.ACTIVE,
        remaining=None,
        type_=None,
    ):
        subscription = Subscription.objects.create(
            organization=self.org,
            child=child or self.child(),
            subscription_type_version=(type_ or self.type).versions.latest(),
            direction=direction or self.ballet,
            branch=branch or self.abaya,
            starts_on=starts_on,
            ends_on=ends_on or starts_on + timedelta(days=30),
            list_price=price,
            price=price,
            renewed_from=renewed_from,
            status=status,
            sessions_remaining_cache=remaining,
        )
        created_on = created_on or min(starts_on, self.today)
        Subscription.objects.filter(pk=subscription.pk).update(
            created_at=self.tz.localize(datetime.combine(created_on, time(12)))
        )
        subscription.refresh_from_db()
        return subscription

    def ended(self, count, *, renewed, ends_on, renewal_price=32000):
        """`count` закончившихся абонементов, из них `renewed` продлены
        следующим абонементом того же направления через 3 дня."""
        for i in range(count):
            old = self.sub(
                starts_on=ends_on - timedelta(days=30),
                ends_on=ends_on,
                status=Subscription.Status.EXPIRED,
            )
            if i < renewed:
                # Продление ещё идёт: заканчивается позже следующего месяца,
                # чтобы само не попасть ни в выборку, ни в прогноз.
                start = ends_on + timedelta(days=3)
                self.sub(
                    child=old.child,
                    starts_on=start,
                    ends_on=self.today + timedelta(days=90),
                    price=renewal_price,
                    created_on=start,
                )

    def history_end(self):
        """Дата окончания, у которой окно продления уже точно прошло."""
        return self.today - timedelta(days=RENEWAL_GRACE_DAYS + 40)

    def forecast(self, branch_ids=None):
        return revenue_forecast(Scope(self.org, branch_ids))


class PrepaidAndDebtTests(ForecastFixtures):
    def test_prepaid_is_paid_minus_worked_off_part(self):
        # 8 занятий за 30 000: осталось 6 — отработано 7 500.
        full = self.sub(starts_on=self.today - timedelta(days=5), remaining=6)
        self.pay(full, 30000)  # 22 500 впереди
        partial = self.sub(starts_on=self.today - timedelta(days=5), remaining=6)
        self.pay(partial, 10000)  # 2 500 впереди, 20 000 долг
        little = self.sub(starts_on=self.today - timedelta(days=5), remaining=6)
        self.pay(little, 5000)  # оплата меньше отработанного: впереди 0, долг 25 000
        future = self.sub(starts_on=self.today + timedelta(days=3))
        self.pay(future, 30000)  # не начат — весь впереди
        expired = self.sub(
            starts_on=self.today - timedelta(days=40),
            status=Subscription.Status.EXPIRED,
            remaining=3,
        )
        self.pay(expired, 30000)  # истёк — центр уже ничего не должен

        data = self.forecast()
        self.assertEqual(data["prepaid"], {"value": Decimal(55000), "subscriptions": 3})
        self.assertEqual(data["unpaid"], {"value": Decimal(45000), "subscriptions": 2})

    def test_unlimited_is_counted_by_days(self):
        unlimited = create_type(
            self.org, name="Безлимит", price=30000, is_unlimited=True, duration_days=30
        )
        sub = self.sub(
            starts_on=self.today - timedelta(days=10),
            ends_on=self.today + timedelta(days=20),
            type_=unlimited,
        )
        self.pay(sub, 30000)
        self.assertEqual(self.forecast()["prepaid"]["value"], Decimal(20000))

    def test_frozen_counts_and_branch_scope(self):
        frozen = self.sub(
            starts_on=self.today - timedelta(days=5),
            remaining=8,
            status=Subscription.Status.FROZEN,
            branch=self.saina,
        )
        self.pay(frozen, 30000)
        self.assertEqual(self.forecast()["prepaid"]["value"], Decimal(30000))
        self.assertEqual(self.forecast((self.abaya.pk,))["prepaid"]["value"], Decimal(0))

    def test_unpaid_is_the_debtors_screen_total(self):
        self.pay(self.sub(starts_on=self.today, branch=self.saina), 10000)
        self.sub(starts_on=self.today - timedelta(days=60), status=Subscription.Status.EXPIRED)
        self.client.force_authenticate(self.owner)
        screen = self.client.get("/api/v1/subscriptions/debtors/").data["total_debt"]
        self.assertEqual(self.forecast()["unpaid"]["value"], Decimal(screen))
        self.assertEqual(self.forecast((self.saina.pk,))["unpaid"]["value"], Decimal(20000))


class RenewalConversionTests(ForecastFixtures):
    def conversion(self, as_of=None):
        index = RenewalIndex(load_rows(self.org, ends_since=add_months(self.today, -12)))
        return renewal_conversion(index, as_of=as_of or self.today, months=6)

    def test_renewal_rules(self):
        ends = self.history_end()
        later = self.today + timedelta(days=90)  # следующие абонементы ещё идут
        # Тот же ребёнок, то же направление, через 3 дня — продлил.
        self.ended(1, renewed=1, ends_on=ends)
        # Продал кнопкой «Продлить» в другое направление — тоже продлил.
        old = self.sub(starts_on=ends - timedelta(days=30), ends_on=ends)
        self.sub(
            child=old.child,
            starts_on=ends + timedelta(days=1),
            ends_on=later,
            direction=self.stretch,
            renewed_from=old,
        )
        # Другое направление без связи — не продление.
        other = self.sub(starts_on=ends - timedelta(days=30), ends_on=ends)
        self.sub(
            child=other.child,
            starts_on=ends + timedelta(days=1),
            ends_on=later,
            direction=self.stretch,
        )
        # Вернулся позже окна продления — не продление.
        late = self.sub(starts_on=ends - timedelta(days=30), ends_on=ends)
        self.sub(
            child=late.child,
            starts_on=ends + timedelta(days=RENEWAL_GRACE_DAYS + 1),
            ends_on=later,
        )

        result = self.conversion()
        self.assertEqual((result["ended"], result["renewed"]), (4, 2))
        self.assertEqual(result["rate"], 0.5)

    def test_window_not_closed_yet_is_not_in_sample(self):
        # Закончился неделю назад — «не продлил» записывать рано.
        self.sub(starts_on=self.today - timedelta(days=37), ends_on=self.today - timedelta(days=7))
        self.assertEqual(self.conversion()["ended"], 0)

    def test_as_of_does_not_see_later_sales(self):
        ends = self.history_end()
        self.ended(1, renewed=1, ends_on=ends)
        as_of = ends + timedelta(days=RENEWAL_GRACE_DAYS + 1)  # продление продано после
        late_look = self.conversion(as_of=as_of + timedelta(days=10))
        early_look = self.conversion(as_of=ends + timedelta(days=2))
        self.assertEqual(late_look["renewed"], 1)
        self.assertEqual(early_look["ended"], 0)  # окно ещё не прошло


class ForecastTests(ForecastFixtures):
    def expiring_next_month(self, count):
        ends = self.next_month + timedelta(days=9)
        return [
            self.sub(starts_on=ends - timedelta(days=30), ends_on=ends, remaining=2)
            for _ in range(count)
        ]

    def test_hidden_when_sample_is_small(self):
        self.ended(MIN_SAMPLE_RANGE - 1, renewed=10, ends_on=self.history_end())
        self.expiring_next_month(5)
        result = self.forecast()["forecast"]
        self.assertEqual(result["status"], "hidden")
        self.assertIsNone(result["value"])
        self.assertIsNone(result["low"])
        self.assertEqual(result["need_more"], 1)
        self.assertEqual(result["expiring"], 5)

    def test_range_only_while_sample_is_medium(self):
        self.ended(30, renewed=20, ends_on=self.history_end())
        self.expiring_next_month(5)
        result = self.forecast()["forecast"]
        self.assertEqual(result["status"], "range")
        self.assertIsNone(result["value"])
        low, high = wilson(20, 30)
        self.assertEqual(result["low"], round(Decimal(str(5 * low)) * 32000))
        self.assertEqual(result["high"], round(Decimal(str(5 * high)) * 32000))
        self.assertLess(result["low"], result["high"])

    def test_point_uses_real_conversion_and_renewal_price(self):
        self.ended(MIN_SAMPLE_POINT, renewed=40, ends_on=self.history_end())
        expiring = self.expiring_next_month(6)
        # Уже продлённый заранее — в прогноз не входит.
        self.sub(child=expiring[0].child, starts_on=expiring[0].ends_on + timedelta(days=1))
        # Истёкший по занятиям — уже не активный, ждать его продления в том месяце нечего.
        Subscription.objects.filter(pk=expiring[1].pk).update(status=Subscription.Status.EXHAUSTED)
        result = self.forecast()["forecast"]
        self.assertEqual(result["status"], "point")
        self.assertEqual(result["expiring"], 4)
        self.assertEqual(result["conversion"]["rate"], 66.7)
        self.assertEqual(result["avg_check"], Decimal(32000))
        # 4 × 40/60 × 32 000
        self.assertEqual(result["value"], Decimal(85333))
        self.assertLessEqual(result["low"], result["value"])
        self.assertGreaterEqual(result["high"], result["value"])

    def test_monthly_renewal_of_renewal_counts(self):
        """Абонемент на 20 дней заканчивается в конце этого месяца: его
        продление закончится в следующем — и его тоже продлят с той же
        вероятностью. Без цепочки такие абонементы из прогноза выпадали бы."""
        self.ended(MIN_SAMPLE_POINT, renewed=40, ends_on=self.history_end())
        self.expiring_next_month(4)
        ends = self.next_month - timedelta(days=1)
        self.sub(starts_on=ends - timedelta(days=20), ends_on=ends, remaining=2)
        result = self.forecast()["forecast"]
        self.assertEqual(result["expiring"], 4)
        rate = 40 / 60
        self.assertEqual(result["expected_renewals"], round(4 * rate + rate**2, 1))
        self.assertEqual(result["value"], round(Decimal(str(4 * rate + rate**2)) * 32000))

    def test_retrospective_compares_forecast_with_fact(self):
        target = add_months(self.month, -1)  # прошлый месяц
        made_on = add_months(target, -1)
        # История к моменту прогноза: 60 закончились за 2 месяца до made_on, 30 продлены.
        old_ends = made_on - timedelta(days=RENEWAL_GRACE_DAYS + 20)
        self.ended(MIN_SAMPLE_POINT, renewed=30, ends_on=old_ends, renewal_price=30000)
        # В прошлом месяце заканчивались 4 абонемента, проданные до made_on.
        ends = target + timedelta(days=4)
        rows = [
            self.sub(
                starts_on=ends - timedelta(days=30),
                ends_on=ends,
                created_on=made_on - timedelta(days=5),
                status=Subscription.Status.EXPIRED,
            )
            for _ in range(4)
        ]
        # Факт: продлили трое, уже после того, как прогноз был сделан.
        for row in rows[:3]:
            self.sub(child=row.child, starts_on=ends + timedelta(days=1), price=36000)

        retro = {r["month"]: r for r in self.forecast()["retrospective"]}
        row = retro[target.isoformat()]
        self.assertEqual(row["as_of"], made_on.isoformat())
        self.assertEqual(row["expiring"], 4)
        self.assertEqual(row["status"], "point")
        self.assertEqual(row["value"], Decimal(60000))  # 4 × 50% × 30 000
        self.assertEqual((row["fact_renewed"], row["fact"]), (3, Decimal(108000)))
        self.assertEqual(row["deviation_percent"], 80.0)
        self.assertFalse(row["in_range"])
        self.assertEqual(len(retro), 6)
        self.assertNotIn(self.month.isoformat(), retro)  # текущий месяц ещё не прошёл


class CalibrationTests(ForecastFixtures):
    def forecast_row(self):
        return {"value": Decimal(100000), "low": Decimal(99000), "high": Decimal(101000)}

    def retro(self, *deviations, complete=True):
        return [{"complete": complete, "deviation_percent": d} for d in deviations]

    def test_range_is_widened_to_worst_past_error(self):
        forecast = self.forecast_row()
        calibrate(forecast, self.retro(-2.0, 5.0, 1.0))
        self.assertEqual((forecast["low"], forecast["high"]), (Decimal(95000), Decimal(105000)))
        self.assertEqual(forecast["calibration_percent"], 5.0)

    def test_not_enough_past_months_or_unfinished_do_not_count(self):
        forecast = self.forecast_row()
        calibrate(forecast, self.retro(-2.0, 30.0) + self.retro(50.0, complete=False))
        self.assertEqual((forecast["low"], forecast["high"]), (Decimal(99000), Decimal(101000)))
        self.assertIsNone(forecast["calibration_percent"])

    def test_hidden_forecast_stays_hidden(self):
        forecast = {"value": None, "low": None, "high": None}
        calibrate(forecast, self.retro(1.0, 2.0, 3.0))
        self.assertIsNone(forecast["low"])


class ForecastApiTests(ForecastFixtures):
    def test_owner_sees_three_separate_values(self):
        self.client.force_authenticate(self.owner)
        response = self.client.get(URL)
        self.assertEqual(response.status_code, 200)
        for key in ("prepaid", "unpaid", "forecast", "retrospective", "rules"):
            self.assertIn(key, response.data)
        self.assertNotIn("_expiring_rows", response.data["forecast"])

    def test_roles(self):
        admin = self.user("+77010000099", User.Role.ADMIN)
        self.client.force_authenticate(admin)
        self.assertEqual(self.client.get(URL).status_code, 403)

    def test_manager_sees_only_own_branch(self):
        self.pay(self.sub(starts_on=self.today, branch=self.saina, remaining=8), 30000)
        manager = self.user("+77010000098", User.Role.MANAGER, branches=[self.abaya])
        self.client.force_authenticate(manager)
        response = self.client.get(URL, {"branch": str(self.saina.pk)})
        self.assertEqual(response.status_code, 403)
        response = self.client.get(URL)
        self.assertEqual(Decimal(str(response.data["prepaid"]["value"])), Decimal(0))

    def test_export(self):
        self.pay(self.sub(starts_on=self.today, remaining=8), 30000)
        self.client.force_authenticate(self.owner)
        response = self.client.get("/api/v1/analytics/export/", {"report": "forecast"})
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames, ["Три величины", "Расчёт прогноза", "Ретроспектива"])
        values = [row[0] for row in book["Три величины"].iter_rows(values_only=True)]
        self.assertIn("Оплачено, но не отработано (обязательства центра)", values)
        self.assertIn("Продано, но не оплачено (задолженность)", values)
        # Разные деньги — без строки «Итого», чтобы их не сложили.
        self.assertNotIn("Итого", values)
        prepaid = next(
            row
            for row in book["Три величины"].iter_rows(values_only=True)
            if row[0] and str(row[0]).startswith("Оплачено")
        )
        self.assertEqual(prepaid[1], 30000)
        retro = [row[0] for row in book["Ретроспектива"].iter_rows(values_only=True)]
        self.assertEqual(len(retro) - retro.index("Месяц") - 1, 6)


@tag("tenant_isolation")
class ForecastTenantIsolationTests(ForecastFixtures):
    def test_other_center_is_not_counted(self):
        other = Organization.objects.create(name="Другой", slug="forecast-other")
        branch = Branch.objects.create(organization=other, name="Чужой")
        direction = Direction.objects.create(organization=other, name="Балет")
        other_type = create_type(other, name="8", price=50000, quota_sessions=8, duration_days=30)
        child = Child.objects.create(
            organization=other, full_name="Чужой", birth_date=self.today, gender="female"
        )
        Subscription.objects.create(
            organization=other,
            child=child,
            subscription_type_version=other_type.versions.latest(),
            direction=direction,
            branch=branch,
            starts_on=self.today,
            ends_on=self.next_month + timedelta(days=3),
            list_price=50000,
            price=50000,
        )
        self.client.force_authenticate(self.owner)
        data = self.client.get(URL).data
        self.assertEqual(Decimal(str(data["unpaid"]["value"])), Decimal(0))
        self.assertEqual(data["forecast"]["expiring"], 0)
        response = self.client.get(URL, {"branch": str(branch.pk)})
        self.assertEqual(response.status_code, 403)
