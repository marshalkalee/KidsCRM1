"""
Выгрузка любого отчёта в Excel (TRU-114, ТЗ раздел 7). Один механизм на
все отчёты: отчёт описывает свои таблицы (`Section`), а шапку, форматы
чисел и итоги с формулами делает этот модуль.

- Выгружается то, что на экране: тот же период, филиалы и фильтры,
  цифры — из тех же функций, что API (compute, breakdown, funnel).
- Числа — числами, даты — датами, деньги и проценты — с форматом ячейки,
  а не текстом: в Excel их можно складывать без ручной правки.
- Шапка каждого листа: отчёт, центр, период, филиалы, фильтры, когда
  выгружено — понятно без открытия системы.
- Под числовыми колонками — строка «Итого» с формулой SUM (там, где сумма
  имеет смысл: не для процентов).

Отчёты выгружают агрегаты — сотни строк, не сотни тысяч, поэтому
синхронно; замер на целевом объёме — в тестах и ADR-0006.
"""

import io
from dataclasses import dataclass, field
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from domains.platform.core.utils import now_for_org

# Формат ячейки по виду значения. Деньги — целые тенге, знак ₸ — в заголовке.
FORMATS = {
    "money": "#,##0",
    "count": "#,##0",
    "percent": "0.0",
    "decimal": "#,##0.0",
    "date": "DD.MM.YYYY",
    "text": "@",
}
SUMMABLE = {"money", "count"}
HEADER_FILL = PatternFill("solid", fgColor="FFE9EC")


@dataclass
class Column:
    title: str
    kind: str = "text"  # text, money, count, percent, decimal, date
    width: int = 16
    # Итог по колонке: None — сумма для денег и штук, False — без итога.
    total: bool | None = None

    @property
    def summable(self):
        return self.kind in SUMMABLE if self.total is None else self.total


@dataclass
class Section:
    """Одна таблица — один лист."""

    title: str
    columns: list
    rows: list = field(default_factory=list)
    note: str = ""


@dataclass
class Export:
    report: str
    sections: list
    filters: list = field(default_factory=list)  # ["Источник: Instagram", …]


def _number(value, kind):
    if value is None or value == "":
        return None
    if kind == "date":
        return date.fromisoformat(value) if isinstance(value, str) else value
    if kind == "text":
        return str(value)
    return float(value) if kind in ("percent", "decimal") else int(round(float(value)))


def _branches_label(scope):
    if scope.branch_ids is None:
        return "Все филиалы"
    return ", ".join(branch.name for branch in scope.branches)


def workbook(export: Export, scope, period) -> bytes:
    book = openpyxl.Workbook()
    book.remove(book.active)
    exported = now_for_org(scope.organization)
    header = [
        (export.report, True),
        (f"Центр: {scope.organization.name}", False),
        (f"Период: {period.start:%d.%m.%Y} – {period.end:%d.%m.%Y}", False),
        (f"Филиалы: {_branches_label(scope)}", False),
        *[(f"Фильтр — {item}", False) for item in export.filters],
        (f"Выгружено: {exported:%d.%m.%Y %H:%M} (время центра)", False),
    ]
    used = set()
    for section in export.sections:
        # Имя листа Excel — до 31 символа и без повторов.
        name = section.title[:31]
        suffix = 2
        while name in used:
            name = f"{section.title[:28]} {suffix}"
            suffix += 1
        used.add(name)
        sheet = book.create_sheet(name)
        for text, bold in header:
            sheet.append([text])
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=bold, size=13 if bold else 11)
        if section.note:
            sheet.append([section.note])
            sheet.cell(row=sheet.max_row, column=1).font = Font(italic=True)
        sheet.append([])
        sheet.append([column.title for column in section.columns])
        # Номер строки — после записи: пустая строка-разделитель в openpyxl
        # сдвигает счётчик, но не max_row, и «+1» промахивался бы на строку.
        head_row = sheet.max_row
        for index, column in enumerate(section.columns, start=1):
            cell = sheet.cell(row=head_row, column=index)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(index)].width = column.width
        for row in section.rows:
            # Ячейка может нести свой вид: (значение, "money") — в таблице
            # «Показатели» в одной колонке и тенге, и штуки, и проценты.
            cells = [
                value if isinstance(value, tuple) else (value, column.kind)
                for value, column in zip(row, section.columns, strict=True)
            ]
            sheet.append([_number(value, kind) for value, kind in cells])
            for index, (_, kind) in enumerate(cells, start=1):
                sheet.cell(row=sheet.max_row, column=index).number_format = FORMATS[kind]
        first, last = head_row + 1, sheet.max_row
        if section.rows and any(column.summable for column in section.columns):
            total_row = last + 1
            sheet.cell(row=total_row, column=1, value="Итого").font = Font(bold=True)
            for index, column in enumerate(section.columns, start=1):
                if index > 1 and column.summable:
                    letter = get_column_letter(index)
                    cell = sheet.cell(
                        row=total_row, column=index, value=f"=SUM({letter}{first}:{letter}{last})"
                    )
                    cell.number_format = FORMATS[column.kind]
                    cell.font = Font(bold=True)
        sheet.freeze_panes = sheet.cell(row=head_row + 1, column=1)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def filename(name, period):
    return f"{name}-{period.start:%Y%m%d}-{period.end:%Y%m%d}.xlsx"
