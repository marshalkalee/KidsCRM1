"""Отток (TRU-127, ТЗ раздел 7)."""

import io
from datetime import date, timedelta
from unittest.mock import patch

import openpyxl
from django.test import tag

from domains.money.subscriptions.models import Subscription
from domains.money.subscriptions.subscription_types import create_type
from domains.people.clients.models import Child, ChildContact, ContactPhone, ParentContact
from domains.platform.tasks.models import Task
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.users.models import User

from .churn import ChurnRules, churn_report
from .period import Period
from .risk_list import risk_list
from .scope import Scope
from .tests_forecast import ForecastFixtures

URL = "/api/v1/analytics/churn/"
EXPORT_URL = "/api/v1/analytics/export/"
SEPTEMBER = Period(date(2025, 9, 1), date(2025, 9, 30), "custom")


class ChurnFixtures(ForecastFixtures):
    today_is = date(2025, 11, 20)

    def setUp(self):
        super().setUp()
        self.set_today(self.today_is)

    def set_today(self, day):
        patcher = patch("domains.platform.analytics.churn.today_for_org", return_value=day)
        patcher.start()
        self.addCleanup(patcher.stop)

    def run_sub(self, child, starts_on, ends_on, **kwargs):
        kwargs.setdefault("status", Subscription.Status.EXPIRED)
        return self.sub(child=child, starts_on=starts_on, ends_on=ends_on, **kwargs)

    def kid(self, name, *, left_reason=None):
        child = self.child(name)
        if left_reason is not None:
            Child.objects.filter(pk=child.pk).update(
                status=Child.Status.LEFT, leave_reason=left_reason
            )
            child.refresh_from_db()
        return child

    def report(self, period=SEPTEMBER, branch_ids=None):
        return churn_report(Scope(self.org, branch_ids), period)

    def names(self, data):
        return sorted(item["name"] for item in data["items"])


class RulesTests(ChurnFixtures):
    def test_deadline_waits_for_september_only_around_summer(self):
        rules = ChurnRules(30, True)
        self.assertEqual(rules.deadline(date(2025, 4, 20)), date(2025, 5, 20))
        self.assertEqual(rules.deadline(date(2025, 5, 5)), date(2025, 9, 30))
        self.assertEqual(rules.deadline(date(2025, 8, 20)), date(2025, 9, 30))
        self.assertEqual(rules.deadline(date(2025, 9, 15)), date(2025, 10, 15))
        self.assertEqual(ChurnRules(30, False).deadline(date(2025, 5, 5)), date(2025, 6, 4))


class ManualCountTests(ChurnFixtures):
    def test_september_matches_manual_count(self):
        """Ручной подсчёт: в сентябре ходили 6 детей, ушли 4."""
        self.run_sub(self.kid("А не вернулась"), date(2025, 8, 11), date(2025, 9, 10))
        b = self.kid("Б продлила через 20 дней")
        self.run_sub(b, date(2025, 8, 16), date(2025, 9, 15))
        self.run_sub(b, date(2025, 10, 5), date(2025, 12, 5), status=Subscription.Status.ACTIVE)
        c = self.kid("В вернулась в ноябре")
        self.run_sub(c, date(2025, 8, 13), date(2025, 9, 12))
        self.run_sub(c, date(2025, 11, 1), date(2025, 12, 1), status=Subscription.Status.ACTIVE)
        d = self.kid("Г переезд", left_reason="Переезд в Астану")
        self.run_sub(d, date(2025, 8, 21), date(2025, 9, 20))
        e = self.kid("Д ходит на растяжку")
        self.run_sub(e, date(2025, 8, 26), date(2025, 9, 25))
        self.run_sub(
            e,
            date(2025, 9, 1),
            date(2025, 12, 1),
            direction=self.stretch,
            status=Subscription.Status.ACTIVE,
        )
        self.run_sub(
            self.kid("Е другой филиал"), date(2025, 8, 9), date(2025, 9, 8), branch=self.saina
        )
        self.run_sub(self.kid("Ж ушла в августе"), date(2025, 7, 1), date(2025, 8, 1))

        data = self.report()
        summary = data["summary"]

        self.assertEqual(
            self.names(data),
            ["А не вернулась", "В вернулась в ноябре", "Г переезд", "Е другой филиал"],
        )
        self.assertEqual(summary["departed"], 4)
        self.assertEqual(summary["returned"], 1)
        self.assertEqual(summary["not_marked"], 3)
        self.assertEqual(summary["active"], 6)
        self.assertEqual(summary["rate"], 66.7)
        self.assertTrue(summary["complete"])
        # Не вернувшиеся — сверху, вернувшаяся — в конце.
        self.assertEqual(data["items"][-1]["returned_on"], "2025-11-01")
        moved = next(item for item in data["items"] if item["marked_left"])
        self.assertEqual(moved["reason"], "Переезд в Астану")
        self.assertEqual(moved["left_on"], "2025-09-20")

        scoped = self.report(branch_ids=(self.abaya.pk,))
        self.assertEqual(scoped["summary"]["departed"], 3)
        self.assertEqual(scoped["summary"]["active"], 5)

    def test_recent_end_is_not_churn_yet(self):
        """Абонемент кончился 10 дней назад — это риск-лист, не отток."""
        self.run_sub(self.kid("Свежая"), date(2025, 10, 11), date(2025, 11, 10))
        marked = self.kid("Отмечена", left_reason="Не понравилось")
        self.run_sub(marked, date(2025, 10, 11), date(2025, 11, 10))

        data = self.report(Period(date(2025, 11, 1), date(2025, 11, 20), "custom"))

        self.assertEqual(self.names(data), ["Отмечена"])
        self.assertEqual(data["summary"]["recent"], 1)
        self.assertFalse(data["summary"]["complete"])

    def test_threshold_comes_from_org_setting(self):
        self.run_sub(self.kid("Ушла"), date(2025, 9, 11), date(2025, 10, 10))
        october = Period(date(2025, 10, 1), date(2025, 10, 31), "custom")
        self.assertEqual(self.report(october)["summary"]["departed"], 1)

        self.org.settings = {**self.org.settings, "churn_inactive_days": 60}
        self.org.save()

        data = self.report(october)
        self.assertEqual(data["summary"]["departed"], 0)
        self.assertEqual(data["summary"]["recent"], 1)
        self.assertEqual(data["rules"]["inactive_days"], 60)


class SummerTests(ChurnFixtures):
    today_is = date(2025, 9, 15)

    def setUp(self):
        super().setUp()
        self.paused = self.kid("Ушла на лето")
        self.run_sub(self.paused, date(2025, 4, 28), date(2025, 5, 28))
        back = self.kid("Вернулась в сентябре")
        self.run_sub(back, date(2025, 5, 16), date(2025, 6, 15))
        self.run_sub(back, date(2025, 9, 5), date(2025, 10, 5), status=Subscription.Status.ACTIVE)
        self.run_sub(self.kid("Ушла в апреле"), date(2025, 3, 21), date(2025, 4, 20))
        self.run_sub(
            self.kid("Отмечена в июне", left_reason="Выросла"), date(2025, 5, 11), date(2025, 6, 10)
        )
        self.spring = Period(date(2025, 4, 1), date(2025, 6, 30), "custom")

    def test_summer_pause_is_not_churn_until_october(self):
        data = self.report(self.spring)

        self.assertEqual(self.names(data), ["Отмечена в июне", "Ушла в апреле"])
        self.assertEqual(data["summary"]["summer_waiting"], 1)
        self.assertEqual(data["summary"]["summer_returned"], 1)
        may = next(row for row in data["trend"] if row["month"] == "2025-05-01")
        self.assertEqual((may["departed"], may["summer_waiting"]), (0, 1))
        self.assertFalse(may["complete"])
        april = next(row for row in data["trend"] if row["month"] == "2025-04-01")
        self.assertTrue(april["complete"])

    def test_after_september_the_unreturned_have_left(self):
        self.set_today(date(2025, 10, 5))

        data = self.report(self.spring)

        self.assertEqual(self.names(data), ["Отмечена в июне", "Ушла в апреле", "Ушла на лето"])
        paused = next(item for item in data["items"] if item["name"] == "Ушла на лето")
        self.assertEqual(paused["left_on"], "2025-05-28")
        may = next(row for row in data["trend"] if row["month"] == "2025-05-01")
        self.assertEqual((may["departed"], may["summer_waiting"], may["complete"]), (1, 0, True))

    def test_without_summer_pause_june_counts_at_once(self):
        self.org.settings = {**self.org.settings, "churn_summer_pause": False}
        self.org.save()

        data = self.report(self.spring)

        self.assertEqual(
            self.names(data),
            ["Вернулась в сентябре", "Отмечена в июне", "Ушла в апреле", "Ушла на лето"],
        )
        self.assertEqual(data["summary"]["summer_waiting"], 0)

    def test_same_month_last_year(self):
        self.run_sub(self.kid("Год назад"), date(2024, 4, 1), date(2024, 4, 30))

        data = self.report(Period(date(2025, 4, 1), date(2025, 4, 30), "custom"))

        self.assertEqual(data["summary"]["previous_year"]["departed"], 1)
        april = data["trend"][-1]
        self.assertEqual(april["month"], "2025-04-01")
        self.assertEqual(april["previous_year"]["month"], "2024-04-01")
        self.assertEqual(april["previous_year"]["departed"], 1)
        self.assertEqual(len(data["trend"]), 12)


class LifetimeAndBreakdownTests(ChurnFixtures):
    def test_lifetime_spans_summer_pause(self):
        long = self.kid("Ходила год")
        self.run_sub(long, date(2024, 1, 1), date(2024, 5, 31))
        self.run_sub(long, date(2024, 9, 10), date(2024, 12, 31))
        self.run_sub(self.kid("Месяц"), date(2024, 10, 1), date(2024, 10, 31))

        data = self.report(Period(date(2024, 10, 1), date(2024, 12, 31), "custom"))

        lifetime = data["lifetime"]["period"]
        self.assertEqual(lifetime["departed"], 2)
        self.assertEqual(lifetime["average_months"], 6.5)  # (366 + 31) / 2 дн.
        buckets = {row["key"]: row["value"] for row in lifetime["buckets"]}
        self.assertEqual(buckets["under_3"], 1)
        self.assertEqual(buckets["12_24"], 1)
        long_item = next(item for item in data["items"] if item["name"] == "Ходила год")
        self.assertEqual(long_item["started_on"], "2024-01-01")
        self.assertEqual(long_item["summer_pauses"], 1)

    def test_reasons_are_grouped_and_unmarked_are_separate(self):
        for name, reason in [("Р1", "Переезд"), ("Р2", " переезд "), ("Р3", "Дорого")]:
            self.run_sub(self.kid(name, left_reason=reason), date(2025, 8, 11), date(2025, 9, 10))
        self.run_sub(self.kid("Без отметки"), date(2025, 8, 11), date(2025, 9, 10))

        reasons = {row["label"]: row["departed"] for row in self.report()["breakdowns"]["reason"]}

        self.assertEqual(reasons, {"Переезд": 2, "Дорого": 1, "Не отмечен ушедшим": 1})

    def test_not_renewed_shows_whether_child_stays(self):
        stays = self.kid("Сменила направление")
        self.run_sub(stays, date(2025, 8, 11), date(2025, 9, 10))
        self.run_sub(
            stays,
            date(2025, 9, 1),
            date(2025, 12, 1),
            direction=self.stretch,
            status=Subscription.Status.ACTIVE,
        )
        self.run_sub(self.kid("Ушла"), date(2025, 8, 11), date(2025, 9, 10))
        renewed = self.kid("Продлила")
        old = self.run_sub(renewed, date(2025, 8, 11), date(2025, 9, 10))
        self.run_sub(renewed, date(2025, 9, 12), date(2025, 10, 12), renewed_from=old)

        states = {item["name"]: item["child_state"] for item in self.report()["not_renewed"]}

        self.assertEqual(states, {"Сменила направление": "active", "Ушла": "departed"})

    def test_contacts_for_call_and_whatsapp(self):
        child = self.kid("С контактом")
        self.run_sub(child, date(2025, 8, 11), date(2025, 9, 10))
        parent = ParentContact.objects.create(
            organization=self.org, full_name="Мама", whatsapp="+77011112233"
        )
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77010000099"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
            is_primary_contact=True,
        )

        item = self.report()["items"][0]

        self.assertEqual(
            (item["parent"], item["phone"], item["whatsapp"]),
            ("Мама", "+77010000099", "+77011112233"),
        )


class RiskListTests(ForecastFixtures):
    """Риск-лист — про тех, кого ещё можно удержать: ушедших в нём нет."""

    def test_departed_are_not_in_risk_list(self):
        gone = self.child("Ушла два месяца назад")
        self.sub(
            child=gone,
            starts_on=self.today - timedelta(days=90),
            ends_on=self.today - timedelta(days=60),
            status=Subscription.Status.EXPIRED,
        )
        fresh = self.child("Истёк неделю назад")
        self.sub(
            child=fresh,
            starts_on=self.today - timedelta(days=37),
            ends_on=self.today - timedelta(days=7),
            status=Subscription.Status.EXPIRED,
        )
        period = Period(self.today - timedelta(days=30), self.today, "custom")

        names = [item["name"] for item in risk_list(Scope(self.org, None), period)["items"]]

        self.assertIn("Истёк неделю назад", names)
        self.assertNotIn("Ушла два месяца назад", names)


class ApiTests(ChurnFixtures):
    params = {"period": "custom", "from": "2025-09-01", "to": "2025-09-30"}

    def test_api_export_and_winback_task(self):
        child = self.kid("Ушла", left_reason="Дорого")
        self.run_sub(child, date(2025, 8, 11), date(2025, 9, 10))
        staying = self.kid("Ходит")
        self.run_sub(staying, date(2025, 9, 1), date(2025, 12, 1), status="active")
        self.client.force_authenticate(self.owner)

        response = self.client.get(URL, self.params)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.data["summary"]["departed"], 1)

        response = self.client.get(EXPORT_URL, {**self.params, "report": "churn"})
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.sheetnames[:4], ["Итог", "Ушли", "Не продлили", "По месяцам"])
        self.assertIn("Причины ухода", book.sheetnames)

        task_url = f"{URL}{child.pk}/task/?period=custom&from=2025-09-01&to=2025-09-30"
        response = self.client.post(task_url)
        self.assertEqual(response.status_code, 201, response.content)
        task = Task.objects.get(pk=response.data["id"])
        self.assertEqual(task.child, child)
        self.assertEqual(task.title, "Вернуть клиента: Ушла")
        self.assertIn("Причина: Дорого", task.description)
        self.assertEqual(self.client.post(task_url).status_code, 200)

        response = self.client.post(
            f"{URL}{staying.pk}/task/?period=custom&from=2025-09-01&to=2025-09-30"
        )
        self.assertEqual(response.status_code, 404)

    def test_teacher_has_no_access(self):
        teacher = self.user("+77010000031", User.Role.TEACHER)
        self.client.force_authenticate(teacher)
        self.assertEqual(self.client.get(URL).status_code, 403)

    def test_settings_api_reads_and_saves_churn_settings(self):
        self.client.force_authenticate(self.owner)
        data = self.client.get("/api/v1/organization/settings/").json()
        self.assertEqual(data["churn_inactive_days"], 30)
        self.assertTrue(data["churn_summer_pause"])

        data.pop("timezones")
        data.pop("public_api_key")
        response = self.client.put(
            "/api/v1/organization/settings/",
            {**data, "churn_inactive_days": 45, "churn_summer_pause": False},
            format="json",
        )
        self.assertEqual(response.status_code, 200, response.content)
        self.org.refresh_from_db()
        self.assertEqual(self.org.settings["churn_inactive_days"], 45)
        self.assertFalse(self.org.settings["churn_summer_pause"])


@tag("tenant_isolation")
class ChurnTenantIsolationTests(ChurnFixtures):
    def test_other_center_is_not_counted(self):
        other = Organization.objects.create(name="Другой", slug="churn-other")
        branch = Branch.objects.create(organization=other, name="Чужой")
        direction = Direction.objects.create(organization=other, name="Балет")
        other_type = create_type(other, name="8", price=50000, quota_sessions=8, duration_days=30)
        child = Child.objects.create(
            organization=other, full_name="Чужой", birth_date=date(2018, 1, 1), gender="female"
        )
        Subscription.objects.create(
            organization=other,
            child=child,
            subscription_type_version=other_type.versions.latest(),
            direction=direction,
            branch=branch,
            starts_on=date(2025, 8, 11),
            ends_on=date(2025, 9, 10),
            list_price=50000,
            price=50000,
            status=Subscription.Status.EXPIRED,
        )
        self.run_sub(self.kid("Своя"), date(2025, 8, 11), date(2025, 9, 10))
        self.client.force_authenticate(self.owner)
        params = {"period": "custom", "from": "2025-09-01", "to": "2025-09-30"}

        data = self.client.get(URL, params).data
        self.assertEqual([item["name"] for item in data["items"]], ["Своя"])

        self.assertEqual(
            self.client.get(URL, {**params, "branch": str(branch.pk)}).status_code, 403
        )
        response = self.client.post(
            f"{URL}{child.pk}/task/?period=custom&from=2025-09-01&to=2025-09-30"
        )
        self.assertEqual(response.status_code, 404)
