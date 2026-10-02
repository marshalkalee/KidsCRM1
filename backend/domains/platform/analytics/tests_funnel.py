"""Воронка продаж по этапам (TRU-115)."""

import io
from datetime import datetime, time, timedelta

import openpyxl
import pytz
from rest_framework.test import APITestCase

from domains.platform.core.utils import today_for_org
from domains.platform.leads.models import Lead, LeadRejectionReason, LeadSource, LeadStatusChange
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.users.models import User

URL = "/api/v1/analytics/funnel/"
S = Lead.Status


class FunnelFixtures(APITestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="tb-funnel")
        self.abaya = Branch.objects.create(organization=self.org, name="Абая")
        self.saina = Branch.objects.create(organization=self.org, name="Саина")
        self.ballet = Direction.objects.create(organization=self.org, name="Балет")
        self.instagram = LeadSource.objects.create(organization=self.org, name="Instagram")
        self.owner = self.user("+77010000001", User.Role.OWNER)
        self.manager = self.user("+77010000002", User.Role.MANAGER, [self.abaya])
        self.today = today_for_org(self.org)
        self.client.force_authenticate(self.owner)
        self.n = 0

    def user(self, phone, role, branches=()):
        user = User.objects.create_user(
            phone=phone, password="x", full_name=phone, organization=self.org, role=role
        )
        user.branches.set(branches)
        return user

    def lead(self, *path, branch=None, source=None, kind="new", days_ago=0):
        """Заявка, прошедшая статусы path по порядку (первый — «новая»)."""
        self.n += 1
        created = pytz.timezone(self.org.timezone).localize(
            datetime.combine(self.today - timedelta(days=days_ago), time(12))
        )
        lead = Lead.objects.create(
            organization=self.org,
            kind=kind,
            parent_name=f"Родитель {self.n}",
            phone=f"+7701{self.n:07d}",
            branch=branch or self.abaya,
            direction=self.ballet,
            source=source,
            assigned_to=self.manager,
            status=path[-1],
        )
        Lead.objects.filter(pk=lead.pk).update(created_at=created)
        previous = ""
        for status in path:
            LeadStatusChange.objects.create(
                organization=self.org, lead=lead, from_status=previous, to_status=status
            )
            previous = status
        return lead


class FunnelTests(FunnelFixtures):
    def stages(self, **params):
        data = self.client.get(URL, params).data["funnel"]
        return data, {stage["key"]: stage["count"] for stage in data["stages"]}

    def test_stages_from_status_history(self):
        self.lead(S.NEW)
        self.lead(S.NEW, S.CONTACTED, S.THINKING)
        self.lead(S.NEW, S.CONTACTED, S.TRIAL_SCHEDULED, S.CONTACTED)  # не дошёл — снова звонок
        self.lead(S.NEW, S.CONTACTED, S.TRIAL_SCHEDULED, S.TRIAL_ATTENDED, S.PURCHASED)
        self.lead(S.NEW, S.CONTACTED, S.PURCHASED)  # купил без пробного
        data, counts = self.stages()
        self.assertEqual(
            counts,
            {"new": 5, "contacted": 4, "trial_scheduled": 2, "trial_attended": 1, "purchased": 2},
        )
        self.assertEqual((data["purchased_after_trial"], data["purchased_without_trial"]), (1, 1))
        self.assertEqual(data["conversion"], 40.0)
        # «Сейчас на этапе» — для клика «застрявшие».
        current = {stage["key"]: stage["current"] for stage in data["stages"]}
        self.assertEqual((current["new"], current["contacted"], current["purchased"]), (1, 1, 2))
        self.assertEqual(data["thinking"], 1)

    def test_renewals_are_not_in_the_funnel(self):
        self.lead(S.NEW, S.CONTACTED, S.PURCHASED, kind="renewal")
        self.lead(S.NEW)
        self.assertEqual(self.stages()[0]["total"], 1)

    def test_rejected_without_contact_is_only_a_lead(self):
        reason = LeadRejectionReason.objects.filter(organization=self.org, kind="new").first()
        lead = self.lead(S.NEW, S.REJECTED)
        Lead.objects.filter(pk=lead.pk).update(rejection_reason=reason)
        data, counts = self.stages()
        self.assertEqual((counts["new"], counts["contacted"], data["rejected"]), (1, 0, 1))

    def test_filters_combine_and_compare(self):
        self.lead(S.NEW, S.CONTACTED, source=self.instagram)
        self.lead(S.NEW, S.CONTACTED, branch=self.saina, source=self.instagram)
        self.lead(S.NEW)
        _, counts = self.stages(source=str(self.instagram.pk), branch=str(self.abaya.pk))
        self.assertEqual((counts["new"], counts["contacted"]), (1, 1))
        # Прошлый месяц — та же когорта за прошлый период.
        self.lead(S.NEW, days_ago=40)
        data, _ = self.stages(
            period="custom",
            **{"from": (self.today - timedelta(days=9)).isoformat(), "to": self.today.isoformat()},
        )
        self.assertEqual(data["total"], 3)
        self.assertIn("previous", data)

    def test_breakdown_by_source_and_manager(self):
        self.lead(S.NEW, S.CONTACTED, S.PURCHASED, source=self.instagram)
        self.lead(S.NEW)
        items = self.client.get(f"{URL}by/", {"by": "source"}).data["items"]
        self.assertEqual(
            [(i["label"], i["total"], i["stages"]["purchased"], i["conversion"]) for i in items],
            [("Instagram", 1, 1, 100.0), (None, 1, 0, 0.0)],
        )
        managers = self.client.get(f"{URL}by/", {"by": "manager"}).data["items"]
        self.assertEqual([(i["label"], i["total"]) for i in managers], [("+77010000002", 2)])
        self.assertEqual(self.client.get(f"{URL}by/", {"by": "weather"}).status_code, 400)

    def test_manager_sees_only_own_branch(self):
        self.lead(S.NEW)
        self.lead(S.NEW, branch=self.saina)
        self.client.force_authenticate(self.manager)
        self.assertEqual(self.stages()[0]["total"], 1)

    def test_excel_export(self):
        """Выгрузка — общим механизмом TRU-114, с фильтром в шапке."""
        self.lead(S.NEW, S.CONTACTED, S.PURCHASED, source=self.instagram)
        self.lead(S.NEW)
        response = self.client.get(
            "/api/v1/analytics/export/", {"report": "funnel", "source": str(self.instagram.pk)}
        )
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(
            book.sheetnames, ["Воронка", "Источник", "Направление", "Филиал", "Ответственный"]
        )
        rows = [row for row in book["Воронка"].iter_rows(values_only=True) if any(row)]
        self.assertIn("Фильтр — Источник: Instagram", [row[0] for row in rows])
        stages = {row[0]: row[1] for row in rows if row[0] in ("Заявки", "Купили")}
        self.assertEqual(stages, {"Заявки": 1, "Купили": 1})

    def test_leads_list_matches_funnel_click(self):
        """Клик «сейчас на этапе» открывает список заявок — те же заявки."""
        self.lead(S.NEW, S.CONTACTED)
        self.lead(S.NEW, S.CONTACTED, days_ago=40)
        response = self.client.get(
            "/api/v1/leads/",
            {
                "status": "contacted",
                "created_from": (self.today - timedelta(days=5)).isoformat(),
                "created_to": self.today.isoformat(),
            },
        )
        self.assertEqual(response.data["count"], 1)
