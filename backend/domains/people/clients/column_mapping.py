"""
Загрузка файла и маппинг колонок (ТЗ п. 4.1) — администратор грузит свой
файл как есть, система угадывает соответствие колонок, остальное
поправляется мышью. Не подменяет предыдущий тикет (шаблон/спецификация,
см. docs/import_format.md) — SYSTEM_FIELDS здесь и колонки шаблона там
должны совпадать по смыслу, просто здесь колонки файла не обязаны
называться ТОЧНО как в шаблоне.

Разбор "грязных" значений одного поля (дата/телефон/роль-префикс) — в
import_service.py (build_row): этот модуль только про то, ОТКУДА взять
сырое значение для каждого поля (какая колонка файла), не про то, как
его почистить.
"""

import csv
import io

import openpyxl

# (ключ_поля, подпись, обязательно) — порядок важен для экрана маппинга.
SYSTEM_FIELDS = [
    ("child_name", "ФИО ребёнка", True),
    ("birth_date", "Дата рождения ребёнка", True),
    ("gender", "Пол ребёнка", True),
    ("parent_name", "ФИО родителя", True),
    ("phone", "Телефон родителя", True),
    ("role", "Роль родителя", False),
    ("medical_notes", "Медицинские заметки", False),
    ("reported_balance", "Остаток занятий", False),
]
REQUIRED_FIELDS = {key for key, _, required in SYSTEM_FIELDS if required}

# Узнаваемые варианты названий колонок для автоподстановки — не только
# точное совпадение с шаблоном (ТЗ: "администратор загружает файл как
# есть"). Ё намеренно продублирована через Е — реальные файлы не
# согласованы в этом (см. _normalize_header).
FIELD_ALIASES = {
    "child_name": [
        "фио ребенка",
        "имя ребенка",
        "ребенок",
        "фамилия имя ребенка",
        "ученик",
        "фио",
    ],
    "birth_date": [
        "дата рождения ребенка",
        "дата рождения",
        "др",
        "дата рожд",
        "birth date",
    ],
    "gender": ["пол ребенка", "пол", "gender"],
    "parent_name": [
        "фио родителя",
        "родитель",
        "имя родителя",
        "контактное лицо",
        "фио контакта",
        "фио мамы",
        "фио папы",
    ],
    "phone": [
        "телефон родителя",
        "телефон",
        "номер телефона",
        "phone",
        "контактный телефон",
        "моб телефон",
        "тел",
    ],
    "role": ["роль родителя", "роль", "кем приходится", "кто"],
    "medical_notes": [
        "медицинские заметки",
        "заметки",
        "медицинская информация",
        "аллергии",
        "здоровье",
    ],
    "reported_balance": [
        "остаток занятий",
        "остаток",
        "баланс",
        "остаток абонемента",
        "осталось занятий",
    ],
}


def _normalize_header(value: str) -> str:
    return " ".join(str(value or "").strip().lower().replace("ё", "е").split())


def guess_mapping(headers: list[str]) -> dict[str, str | None]:
    """Каждому полю системы — заголовок файла или None, если не угадали.
    Точное совпадение с алиасом приоритетнее частичного — иначе короткий
    алиас вроде "тел" мог бы перехватить не ту колонку раньше более
    точного варианта."""
    normalized = {_normalize_header(h): h for h in headers if _normalize_header(h)}
    mapping: dict[str, str | None] = {key: None for key, _, _ in SYSTEM_FIELDS}
    used_headers: set[str] = set()

    for field_key, _, _ in SYSTEM_FIELDS:
        for alias in FIELD_ALIASES.get(field_key, []):
            header = normalized.get(alias)
            if header and header not in used_headers:
                mapping[field_key] = header
                used_headers.add(header)
                break

    for field_key, _, _ in SYSTEM_FIELDS:
        if mapping[field_key] is not None:
            continue
        for norm_header, header in normalized.items():
            if header in used_headers:
                continue
            if any(alias in norm_header for alias in FIELD_ALIASES.get(field_key, [])):
                mapping[field_key] = header
                used_headers.add(header)
                break

    return mapping


def _build_merge_lookup(sheet) -> dict[tuple[int, int], object]:
    """(строка, колонка, 1-based) -> значение якоря объединённой ячейки —
    см. docs/import_format.md, находка про объединённые ячейки родителя
    на несколько строк его детей."""
    lookup = {}
    for merged_range in sheet.merged_cells.ranges:
        anchor = sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                if (row, col) != (merged_range.min_row, merged_range.min_col):
                    lookup[(row, col)] = anchor
    return lookup


def _clean(value) -> str:
    return str(value).strip() if value is not None else ""


def read_xlsx(file) -> tuple[list[str], list[tuple[int, list]]]:
    """read_only=False — нужен доступ к merged_cells.ranges/произвольным
    ячейкам; для разового admin-действия на файл в тысячи строк это не
    проблема (не сравнимо с бюджетом списка на 5000 детей)."""
    workbook = openpyxl.load_workbook(file, data_only=True, read_only=False)
    sheet = workbook.active
    merge_lookup = _build_merge_lookup(sheet)

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [_clean(c) for c in header_row]

    rows = []
    for row_number, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(_clean(v) for v in raw_row):
            continue  # пустая строка-разделитель — не ошибка
        resolved = list(raw_row)
        for idx, value in enumerate(raw_row):
            if value is None:
                resolved[idx] = merge_lookup.get((row_number, idx + 1))
        rows.append((row_number, resolved))
    return headers, rows


# cp1251 — норма для CSV, выгруженного из Excel на Windows (ТЗ), не
# исключение; utf-8-sig первым — так же часто встречается и однозначно
# определяется по BOM.
CSV_ENCODINGS = ("utf-8-sig", "utf-8", "cp1251")


def _decode_csv_bytes(raw_bytes: bytes) -> tuple[str, str]:
    for encoding in CSV_ENCODINGS:
        try:
            return raw_bytes.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw_bytes.decode("utf-8", errors="replace"), "utf-8 (с заменой нечитаемых символов)"


def _detect_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
    except csv.Error:
        # ";" — норма для выгрузки из Excel на Windows (ТЗ), запасной
        # вариант — запятая.
        return ";" if sample.count(";") >= sample.count(",") else ","


def read_csv(file) -> tuple[list[str], list[tuple[int, list]], dict]:
    text, encoding = _decode_csv_bytes(file.read())
    delimiter = _detect_delimiter(text[:2000])
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    all_rows = list(reader)

    if not all_rows:
        return [], [], {"encoding": encoding, "delimiter": delimiter}

    headers = [_clean(c) for c in all_rows[0]]
    rows = []
    for row_number, raw_row in enumerate(all_rows[1:], start=2):
        if not any(_clean(v) for v in raw_row):
            continue
        rows.append((row_number, raw_row))
    return headers, rows, {"encoding": encoding, "delimiter": delimiter}


def read_uploaded_file(file, filename: str) -> tuple[list[str], list[tuple[int, list]], dict]:
    if filename.lower().endswith(".csv"):
        return read_csv(file)
    headers, rows = read_xlsx(file)
    return headers, rows, {}


def apply_mapping(
    headers: list[str], raw_rows: list[tuple[int, list]], mapping: dict[str, str | None]
) -> list[tuple[int, dict]]:
    """Возвращает [(номер_строки, {ключ_поля: сырое_значение})] — сама
    очистка/валидация значений — build_row() в import_service.py."""
    header_index = {h: i for i, h in enumerate(headers)}
    resolved = []
    for row_number, raw_row in raw_rows:
        values = {}
        for field_key, header_name in mapping.items():
            if not header_name:
                continue
            idx = header_index.get(header_name)
            values[field_key] = raw_row[idx] if idx is not None and idx < len(raw_row) else None
        resolved.append((row_number, values))
    return resolved
