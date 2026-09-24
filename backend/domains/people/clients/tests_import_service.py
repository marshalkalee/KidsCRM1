"""
Импорт из Excel (ТЗ п. 4.1, MVP критерий приёмки №1) — разбор файла,
дедуп внутри файла и против базы через ChildService, выполнение.

Разбор файла здесь идёт через полный конвейер column_mapping.py
(read_xlsx/guess_mapping/apply_mapping) + import_service.build_rows —
так же, как это делают реальные веб-экраны (import_views.py), а не
напрямую через build_row с ручным словарём: колонки в HEADERS ниже
специально названы так, чтобы guess_mapping их узнал сама, без ручной
правки маппинга в тесте.
"""

import datetime
import io

import openpyxl
from django.test import TestCase

from domains.money.subscriptions.models import Subscription
from domains.platform.tenants.models import Organization

from . import column_mapping
from .import_service import (
    ImportRow,
    ReportLevel,
    RowAction,
    build_dry_run_report,
    build_rows,
    execute_import,
    resolve_rows,
)
from .models import Child, ChildContact, ContactPhone, ParentContact
from .services import DuplicateReason

HEADERS = [
    "ФИО ребёнка",
    "Дата рождения ребёнка",
    "Пол ребёнка",
    "ФИО родителя",
    "Телефон родителя",
    "Роль родителя",
]


def _workbook(rows, headers=HEADERS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _parse(file):
    headers, raw_rows = column_mapping.read_xlsx(file)
    mapping = column_mapping.guess_mapping(headers)
    mapped_rows = column_mapping.apply_mapping(headers, raw_rows, mapping)
    return build_rows(mapped_rows)


class ParseWorkbookTests(TestCase):
    def test_parses_valid_rows(self):
        file = _workbook(
            [
                [
                    "Айгерим Серикова",
                    "10.03.2018",
                    "ж",
                    "Иванова Марина",
                    "+7 701 123 45 67",
                    "мама",
                ],
            ]
        )

        rows = _parse(file)

        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertTrue(row.is_valid)
        self.assertEqual(row.child_name, "Айгерим Серикова")
        self.assertEqual(row.birth_date, datetime.date(2018, 3, 10))
        self.assertEqual(row.gender, Child.Gender.FEMALE)
        self.assertEqual(row.parent_name, "Иванова Марина")
        self.assertEqual(row.phone, "+77011234567")
        self.assertEqual(row.role, ChildContact.Role.MOTHER)

    def test_missing_required_header_leaves_field_unmapped_and_row_invalid(self):
        # Не "ошибка заголовков" как раньше — колонки без пары просто не
        # попадают в маппинг (см. column_mapping.guess_mapping), и
        # соответствующее поле становится обычной ошибкой строки.
        headers = [h for h in HEADERS if h != "Телефон родителя"]
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "мама"]], headers=headers)

        mapping = column_mapping.guess_mapping(headers)
        self.assertIsNone(mapping["phone"])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertTrue(any("телефон" in e for e in rows[0].errors))

    def test_missing_child_name_is_a_row_error(self):
        file = _workbook([["", "10.03.2018", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertIn("не заполнено ФИО ребёнка", rows[0].errors[0])

    def test_unparseable_birth_date_is_a_row_error(self):
        file = _workbook([["Данияр", "не дата", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertTrue(any("дат" in e for e in rows[0].errors))

    def test_unrecognized_gender_is_a_row_error(self):
        file = _workbook([["Данияр", "10.03.2018", "ктотоеще", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertIn("не удалось разобрать пол ребёнка: ктотоеще (ожидается м/ж)", rows[0].errors)

    def test_blank_gender_error_does_not_show_empty_quotes(self):
        file = _workbook([["Данияр", "10.03.2018", "", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertIn("не указан пол ребёнка", rows[0].errors)

    def test_missing_parent_name_is_a_row_error(self):
        file = _workbook([["Данияр", "10.03.2018", "м", "", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)

    def test_unparseable_phone_is_a_row_error(self):
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "+7 701 12-34", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        # Значение как в файле, без кавычек-repr — понятно администратору (ТЗ п. 10.4).
        self.assertIn("не удалось разобрать телефон: +7 701 12-34", rows[0].errors)

    def test_blank_phone_is_a_row_error(self):
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "", ""]])

        rows = _parse(file)

        self.assertIn("не заполнен телефон", rows[0].errors)

    def test_blank_role_defaults_to_other(self):
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertEqual(rows[0].role, ChildContact.Role.OTHER)

    def test_trailing_blank_row_is_skipped(self):
        file = _workbook(
            [
                ["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", ""],
                [None, None, None, None, None, None],
            ]
        )

        rows = _parse(file)

        self.assertEqual(len(rows), 1)

    def test_excel_native_date_cell_is_parsed(self):
        file = _workbook(
            [["Данияр", datetime.date(2018, 3, 10), "м", "Иванова", "+77011234567", ""]]
        )

        rows = _parse(file)

        self.assertEqual(rows[0].birth_date, datetime.date(2018, 3, 10))

    def test_slash_separated_date_is_parsed(self):
        file = _workbook([["Данияр", "10/03/2018", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertEqual(rows[0].birth_date, datetime.date(2018, 3, 10))

    def test_iso_date_is_parsed(self):
        file = _workbook([["Данияр", "2018-03-10", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertEqual(rows[0].birth_date, datetime.date(2018, 3, 10))

    def test_role_prefix_in_parent_name_is_split_out(self):
        # Реальная грязь из ТЗ: "мама Айгерим" вместо чистого ФИО.
        file = _workbook([["Данияр", "10.03.2018", "м", "мама Иванова Марина", "+77011234567", ""]])

        rows = _parse(file)

        self.assertEqual(rows[0].parent_name, "Иванова Марина")
        self.assertEqual(rows[0].role, ChildContact.Role.MOTHER)

    def test_explicit_role_column_wins_over_name_prefix(self):
        file = _workbook(
            [["Данияр", "10.03.2018", "м", "мама Иванова Марина", "+77011234567", "папа"]]
        )

        rows = _parse(file)

        self.assertEqual(rows[0].role, ChildContact.Role.FATHER)

    def test_multiple_phones_in_one_cell_are_all_captured(self):
        file = _workbook(
            [
                [
                    "Данияр",
                    "10.03.2018",
                    "м",
                    "Иванова Марина",
                    "+7 701 123 45 67, +7 707 890 89 89",
                    "",
                ]
            ]
        )

        rows = _parse(file)

        self.assertEqual(rows[0].phone, "+77011234567")
        self.assertEqual(rows[0].extra_phones, ["+77078908989"])

    def test_one_invalid_phone_among_several_is_dropped_not_fatal(self):
        file = _workbook(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "не телефон, +77011234567", ""]]
        )

        rows = _parse(file)

        self.assertTrue(rows[0].is_valid)
        self.assertEqual(rows[0].phone, "+77011234567")

    def test_medical_notes_column_is_captured(self):
        headers = [*HEADERS, "Медицинские заметки"]
        file = _workbook(
            [["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", "", "Аллергия на орехи"]],
            headers=headers,
        )

        rows = _parse(file)

        self.assertEqual(rows[0].medical_notes, "Аллергия на орехи")

    def test_balance_column_is_captured_not_dropped(self):
        headers = [*HEADERS, "Остаток занятий"]
        file = _workbook(
            [["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", "", "5"]], headers=headers
        )

        rows = _parse(file)

        self.assertEqual(rows[0].reported_balance, "5")

    def test_merged_parent_cells_are_resolved_for_second_child_row(self):
        # Реальная грязь из ТЗ: ФИО/телефон родителя объединены на
        # несколько строк, когда у него больше одного ребёнка в файле.
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"])
        ws.append(["Айгерим", "01.01.2020", "ж", None, None, None])
        ws.merge_cells("D2:D3")
        ws.merge_cells("E2:E3")
        ws.merge_cells("F2:F3")
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        rows = _parse(buf)

        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1].parent_name, "Иванова Марина")
        self.assertEqual(rows[1].phone, "+77011234567")
        self.assertEqual(rows[1].role, ChildContact.Role.MOTHER)


class BirthDatePlausibilityTests(TestCase):
    """Тикет «валидация, сухой прогон и отчёт об ошибках» (ТЗ п. 4.1):
    дата рождения должна не только разобраться, но и быть правдоподобной
    (не 1905 год, не в будущем)."""

    def test_future_birth_date_is_an_error(self):
        future_year = datetime.date.today().year + 1
        file = _workbook([["Данияр", f"01.01.{future_year}", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertTrue(any("будущем" in e for e in rows[0].errors))

    def test_implausibly_old_birth_date_is_an_error(self):
        # Тот самый пример из тикета — "1905" вместо "2005".
        file = _workbook([["Данияр", "01.01.1905", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertFalse(rows[0].is_valid)
        self.assertTrue(any("опечатку" in e for e in rows[0].errors))

    def test_realistic_birth_date_has_no_plausibility_error(self):
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertTrue(rows[0].is_valid)


class DirectionAndGroupWarningTests(TestCase):
    """Тикет «валидация, сухой прогон и отчёт об ошибках» (ТЗ п. 4.1):
    направление/группа существуют — или об этом сказано явно
    (предупреждение, не ошибка: строку не блокирует)."""

    def _row_with_direction_and_group(self, known_directions=frozenset(), known_groups=frozenset()):
        headers = [*HEADERS, "Направление", "Группа"]
        file = _workbook(
            [["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", "", "Балет", "Балет-мини"]],
            headers=headers,
        )
        file_headers, raw_rows = column_mapping.read_xlsx(file)
        mapping = column_mapping.guess_mapping(file_headers)
        mapped_rows = column_mapping.apply_mapping(file_headers, raw_rows, mapping)
        return build_rows(mapped_rows, known_directions, known_groups)[0]

    def test_unknown_direction_and_group_produce_warnings_not_errors(self):
        row = self._row_with_direction_and_group()

        self.assertTrue(row.is_valid)
        self.assertTrue(any("Балет" in w and "направление" in w for w in row.warnings))
        self.assertTrue(any("Балет-мини" in w and "группа" in w for w in row.warnings))

    def test_known_direction_and_group_produce_no_warnings(self):
        row = self._row_with_direction_and_group(
            known_directions=frozenset({"балет"}), known_groups=frozenset({"балет-мини"})
        )

        self.assertEqual(row.warnings, [])

    def test_direction_match_is_case_and_whitespace_insensitive(self):
        row = self._row_with_direction_and_group(
            known_directions=frozenset({"балет"}), known_groups=frozenset({"балет-мини"})
        )
        self.assertEqual(row.direction_name, "Балет")
        self.assertEqual(row.group_name, "Балет-мини")

    def test_empty_direction_and_group_produce_no_warnings(self):
        # Необязательные поля — пустая колонка не повод для предупреждения.
        file = _workbook([["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", ""]])

        rows = _parse(file)

        self.assertEqual(rows[0].warnings, [])


class ResolveRowsTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")

    def _row(self, row_number, name, phone, **extra):
        defaults = dict(
            row_number=row_number,
            child_name=name,
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.FEMALE,
            parent_name="Иванова Марина",
            phone=phone,
            role=ChildContact.Role.MOTHER,
        )
        defaults.update(extra)
        return ImportRow(**defaults)

    def test_row_with_errors_gets_error_action(self):
        row = self._row(2, "Данияр", "+77011234567", errors=["что-то не так"])

        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.ERROR)

    def test_new_row_defaults_to_create_new_family(self):
        row = self._row(2, "Данияр", "+77011234567")

        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.CREATE_NEW_FAMILY)

    def test_two_rows_same_phone_different_children_in_file_are_second_child_not_duplicate(self):
        # Ключевой сценарий тикета внутри одного файла: тот же телефон,
        # другой ребёнок — второй ребёнок в семье, не дубль.
        row1 = self._row(2, "Данияр", "+77011234567")
        row2 = self._row(3, "Айгерим Серикова", "+77011234567")

        resolve_rows(self.org, [row1, row2])

        self.assertEqual(row1.action, RowAction.CREATE_NEW_FAMILY)
        self.assertEqual(row2.action, RowAction.ATTACH_EXISTING)
        self.assertEqual(row2.reason, DuplicateReason.EXISTING_PARENT_NEW_CHILD)
        self.assertIsNone(row2.matched_parent_id)
        self.assertEqual(row2.attached_to_row_number, 2)

    def test_same_phone_same_child_in_file_is_a_duplicate_not_a_second_child(self):
        row1 = self._row(2, "Данияр", "+77011234567")
        row2 = self._row(3, "Данияр", "+77011234567")

        resolve_rows(self.org, [row1, row2])

        # Второй ряд — буквально та же строка (тот же телефон И тот же
        # ребёнок) — не второй ребёнок, а дубль/повтор строки.
        self.assertEqual(row2.action, RowAction.SKIP)
        self.assertEqual(row2.reason, DuplicateReason.PHONE)
        self.assertTrue(any("дубликат" in w and "файл" in w for w in row2.warnings))

    def test_row_matches_existing_db_family_with_different_child(self):
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2016, 1, 1),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")

        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )

        row = self._row(2, "Айгерим Серикова", "+77011234567")
        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.ATTACH_EXISTING)
        self.assertEqual(row.matched_parent_id, str(parent.id))

    def test_second_file_row_still_recognizes_existing_db_child_as_duplicate(self):
        # Найдено живой проверкой: первый ряд файла с телефоном
        # "запоминал" семью, и второй ряд с ТЕМ ЖЕ телефоном сверялся
        # только со строками файла, забывая, что у найденной семьи уже
        # есть свой ребёнок в базе — так реального дубля второй ряд не
        # находил вообще.
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2016, 1, 1),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )

        # Первый ряд файла с этим телефоном — другой ребёнок, реально
        # второй в семье.
        row1 = self._row(2, "Айгерим Серикова", "+77011234567")
        # Второй ряд файла с тем же телефоном — на самом деле уже
        # существующий в базе "Данияр", а не третий ребёнок.
        row2 = self._row(3, "Данияр", "+77011234567", birth_date=datetime.date(2016, 1, 1))

        resolve_rows(self.org, [row1, row2])

        self.assertEqual(row1.action, RowAction.ATTACH_EXISTING)
        self.assertEqual(row2.action, RowAction.SKIP)
        self.assertEqual(row2.reason, DuplicateReason.PHONE)

    def test_row_matching_existing_child_exactly_is_skipped_by_default(self):
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")

        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )

        row = self._row(2, "Данияр", "+77011234567", birth_date=datetime.date(2018, 3, 10))
        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.SKIP)
        self.assertEqual(row.reason, DuplicateReason.PHONE)
        self.assertEqual(row.matched_child.id, existing_child.id)
        self.assertTrue(any("дубликат" in w and "базе" in w for w in row.warnings))

    def test_weak_name_match_produces_a_warning_not_a_block(self):
        # Слабое совпадение из тикета "валидация, сухой прогон и отчёт об
        # ошибках": ФИО похоже, но телефон другой — не блокирует строку.
        Child.objects.create(
            organization=self.org,
            full_name="Данияр Серикулы",
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.MALE,
        )

        row = self._row(2, "Данияр Серикулы", "+77079998877")
        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.CREATE_NEW_FAMILY)
        self.assertTrue(row.is_valid)
        self.assertTrue(any("похоже на уже существующего" in w for w in row.warnings))

    def test_attaching_second_child_produces_no_warning(self):
        # Второй ребёнок в семье — ожидаемый штатный путь, не повод для
        # предупреждения (в отличие от дублей/слабых совпадений).
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2016, 1, 1),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )

        row = self._row(2, "Айгерим Серикова", "+77011234567")
        resolve_rows(self.org, [row])

        self.assertEqual(row.action, RowAction.ATTACH_EXISTING)
        self.assertEqual(row.warnings, [])


class ExecuteImportTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")

    def _row(self, row_number, name, phone, **extra):
        defaults = dict(
            row_number=row_number,
            child_name=name,
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.FEMALE,
            parent_name="Иванова Марина",
            phone=phone,
            role=ChildContact.Role.MOTHER,
        )
        defaults.update(extra)
        return ImportRow(**defaults)

    def test_creates_child_and_new_parent(self):
        row = self._row(2, "Данияр", "+77011234567")
        resolve_rows(self.org, [row])

        result = execute_import(self.org, [row])

        self.assertEqual(result.created, 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)
        self.assertEqual(ParentContact.objects.for_tenant(self.org).count(), 1)

    def test_second_child_same_phone_attaches_to_the_same_parent_not_a_new_one(self):
        row1 = self._row(2, "Данияр", "+77011234567")
        row2 = self._row(3, "Айгерим Серикова", "+77011234567")
        resolve_rows(self.org, [row1, row2])

        result = execute_import(self.org, [row1, row2])

        self.assertEqual(result.created, 1)
        self.assertEqual(result.attached_to_existing_family, 1)
        self.assertEqual(ParentContact.objects.for_tenant(self.org).count(), 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 2)

    def test_skipped_rows_are_not_created(self):
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")

        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )
        row = self._row(2, "Данияр", "+77011234567", birth_date=datetime.date(2018, 3, 10))
        resolve_rows(self.org, [row])

        result = execute_import(self.org, [row])

        self.assertEqual(result.skipped, 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

    def test_error_rows_are_skipped_and_not_created(self):
        row = self._row(2, "Данияр", "+77011234567", errors=["не заполнено ФИО родителя"])
        resolve_rows(self.org, [row])

        result = execute_import(self.org, [row])

        self.assertEqual(result.skipped, 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 0)

    def test_no_automatic_merge_ever_happens(self):
        # Явный критерий приёмки: сервис никогда не объединяет/удаляет
        # существующие записи сам — только создаёт новые или пропускает.
        existing_child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")

        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org,
            child=existing_child,
            parent_contact=parent,
            role=ChildContact.Role.MOTHER,
        )
        row = self._row(2, "Данияр", "+77011234567", birth_date=datetime.date(2018, 3, 10))
        resolve_rows(self.org, [row])
        execute_import(self.org, [row])

        self.assertTrue(Child.objects.filter(pk=existing_child.pk).exists())
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

    def test_extra_phones_are_all_saved_on_new_parent(self):
        row = self._row(2, "Данияр", "+77011234567", extra_phones=["+77078908989"])
        resolve_rows(self.org, [row])

        execute_import(self.org, [row])

        parent = ParentContact.objects.for_tenant(self.org).get()
        self.assertEqual(
            set(parent.phones.values_list("number", flat=True)),
            {"+77011234567", "+77078908989"},
        )

    def test_medical_notes_are_saved_on_the_child(self):
        row = self._row(2, "Данияр", "+77011234567", medical_notes="Аллергия на орехи")
        resolve_rows(self.org, [row])

        execute_import(self.org, [row])

        child = Child.objects.for_tenant(self.org).get()
        self.assertEqual(child.medical_notes, "Аллергия на орехи")

    def test_reported_balance_is_surfaced_not_imported_as_subscription(self):
        row = self._row(2, "Данияр", "+77011234567", reported_balance="5")
        resolve_rows(self.org, [row])

        result = execute_import(self.org, [row])

        self.assertEqual(result.unhandled_balances, [(2, "Данияр", "5")])
        self.assertEqual(Subscription.objects.for_tenant(self.org).count(), 0)


class BuildDryRunReportTests(TestCase):
    """Отчёт сухого прогона (ТЗ п. 4.1): готово/предупреждение/ошибка по
    каждой строке, номер строки — из исходного файла (ТЗ п. 10.4)."""

    def test_error_row_is_counted_as_error_with_its_messages(self):
        row = ImportRow(row_number=147, child_name="", errors=["не заполнено ФИО ребёнка"])

        report = build_dry_run_report([row])

        self.assertEqual(report.error_count, 1)
        self.assertEqual(report.ready_count, 0)
        self.assertEqual(report.warning_count, 0)
        self.assertEqual(
            report.rows,
            [
                {
                    "row_number": 147,
                    "level": ReportLevel.ERROR,
                    "child_name": "",
                    "messages": ["не заполнено ФИО ребёнка"],
                }
            ],
        )

    def test_warning_row_is_counted_as_warning_not_error(self):
        row = ImportRow(row_number=5, child_name="Данияр", warnings=["дубликат — уже есть в базе"])

        report = build_dry_run_report([row])

        self.assertEqual(report.warning_count, 1)
        self.assertEqual(report.error_count, 0)
        self.assertEqual(report.rows[0]["level"], ReportLevel.WARNING)

    def test_clean_row_is_counted_as_ready(self):
        row = ImportRow(row_number=5, child_name="Данияр")

        report = build_dry_run_report([row])

        self.assertEqual(report.ready_count, 1)
        self.assertEqual(report.rows[0]["level"], ReportLevel.READY)
        self.assertEqual(report.rows[0]["messages"], [])

    def test_row_numbers_in_report_are_original_file_rows_not_index(self):
        # ТЗ п. 10.4 — администратор ищет проблему в исходном Excel по
        # номеру строки, не по порядковому индексу в списке.
        rows = [
            ImportRow(row_number=2, child_name="Один"),
            ImportRow(row_number=147, child_name="", errors=["не заполнено ФИО ребёнка"]),
        ]

        report = build_dry_run_report(rows)

        self.assertEqual([r["row_number"] for r in report.rows], [2, 147])
