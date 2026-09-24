"""
Веб-экраны импорта (import_views.py) через настоящий HTTP-клиент
(multipart-загрузка файла), не только юниты сервисов (см.
tests_import_service.py/tests_column_mapping.py/tests_tasks.py).
Шаги: загрузка -> маппинг -> сухой прогон (фоновая задача) -> отчёт ->
запуск настоящего импорта из отчёта (ещё одна фоновая задача) -> статус
(ТЗ п. 4.1, п. 10.1).
"""

import datetime
import io
from unittest import mock

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from domains.platform.tenants.models import Organization

from .import_service import Decision, DuplicateKind
from .models import (
    Child,
    ChildContact,
    CommunicationLog,
    ContactPhone,
    ImportColumnMapping,
    ImportJob,
    ParentContact,
)
from .tasks import run_dry_run_job, run_import_job

User = get_user_model()

HEADERS = [
    "ФИО ребёнка",
    "Дата рождения ребёнка",
    "Пол ребёнка",
    "ФИО родителя",
    "Телефон родителя",
    "Роль родителя",
]


def _xlsx_file(rows, headers=HEADERS, filename="import.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return SimpleUploadedFile(
        filename,
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _mapping_dict(mapping_rows):
    return {key: current for key, _, _, current in mapping_rows}


def _run_celery_tasks_synchronously():
    """.delay() требует брокера (Redis) — в тестах обе задачи выполняются
    синхронно в процессе через apply(), без брокера и без изменения кода
    вьюх/задач (сам run_dry_run_job/run_import_job — см. tests_tasks.py)."""
    patchers = [
        mock.patch.object(
            run_dry_run_job,
            "delay",
            side_effect=lambda job_id: run_dry_run_job.apply(args=[job_id]),
        ),
        mock.patch.object(
            run_import_job, "delay", side_effect=lambda job_id: run_import_job.apply(args=[job_id])
        ),
    ]
    for patcher in patchers:
        patcher.start()
    return patchers


class ChildImportUploadViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.teacher = User.objects.create_user(
            phone="+77010000002",
            full_name="Teacher",
            password="pass12345",
            organization=self.org,
            role=User.Role.TEACHER,
        )

    def test_get_renders_upload_form(self):
        self.client.force_login(self.owner)

        response = self.client.get(reverse("clients_web:child-import-upload"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'enctype="multipart/form-data"')

    def test_teacher_cannot_access_upload(self):
        self.client.force_login(self.teacher)

        response = self.client.get(reverse("clients_web:child-import-upload"))

        self.assertEqual(response.status_code, 403)

    def test_rejects_unsupported_extension(self):
        self.client.force_login(self.owner)
        bad_file = SimpleUploadedFile("import.txt", b"not a spreadsheet", content_type="text/plain")

        response = self.client.post(reverse("clients_web:child-import-upload"), {"file": bad_file})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "xlsx")

    def test_valid_xlsx_upload_shows_mapping_screen_with_guessed_columns(self):
        self.client.force_login(self.owner)
        file = _xlsx_file([["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]])

        response = self.client.post(reverse("clients_web:child-import-upload"), {"file": file})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Данияр")
        mapping = _mapping_dict(response.context["mapping_rows"])
        self.assertEqual(mapping["child_name"], "ФИО ребёнка")
        self.assertEqual(mapping["phone"], "Телефон родителя")

    def test_valid_csv_upload_detects_cp1251_and_semicolon(self):
        # ТЗ явно: файлы из Excel в Windows приходят в cp1251 с точкой
        # с запятой — норма, не исключение.
        self.client.force_login(self.owner)
        text = (
            "ФИО ребёнка;Дата рождения ребёнка;Пол ребёнка;ФИО родителя;"
            "Телефон родителя;Роль родителя\n"
            "Данияр;10.03.2018;м;Иванова Марина;+77011234567;мама\n"
        )
        file = SimpleUploadedFile("import.csv", text.encode("cp1251"), content_type="text/csv")

        response = self.client.post(reverse("clients_web:child-import-upload"), {"file": file})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "cp1251")
        self.assertContains(response, "Данияр")

    def test_empty_file_shows_form_error(self):
        self.client.force_login(self.owner)
        file = _xlsx_file([])

        response = self.client.post(reverse("clients_web:child-import-upload"), {"file": file})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "не нашлось")

    def test_saved_mapping_is_reused_on_next_upload_of_same_columns(self):
        # Критерий приёмки: повторный импорт с тем же составом колонок
        # не требует настраивать маппинг заново.
        self.client.force_login(self.owner)
        custom_headers = ["Child", "DOB", "Sex", "Parent", "Contact Data", "Relation"]
        file = _xlsx_file(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]],
            headers=custom_headers,
        )

        first_upload = self.client.post(reverse("clients_web:child-import-upload"), {"file": file})
        first_mapping = _mapping_dict(first_upload.context["mapping_rows"])
        self.assertIsNone(first_mapping["phone"])  # "Contact Data" угадать нечем

        confirm_data = {
            "headers_json": first_upload.context["headers_json"],
            "raw_rows_json": first_upload.context["raw_rows_json"],
            "meta_json": first_upload.context["meta_json"],
            "mapping_child_name": "Child",
            "mapping_birth_date": "DOB",
            "mapping_gender": "Sex",
            "mapping_parent_name": "Parent",
            "mapping_phone": "Contact Data",
            "mapping_role": "Relation",
        }
        with mock.patch.object(run_dry_run_job, "delay"):
            self.client.post(reverse("clients_web:child-import-mapping-confirm"), confirm_data)

        second_file = _xlsx_file(
            [["Айгерим", "01.01.2020", "ж", "Иванова Марина", "+77011234567", "мама"]],
            headers=custom_headers,
        )
        second_upload = self.client.post(
            reverse("clients_web:child-import-upload"), {"file": second_file}
        )

        second_mapping = _mapping_dict(second_upload.context["mapping_rows"])
        self.assertEqual(second_mapping["phone"], "Contact Data")


class ChildImportMappingConfirmViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)
        for patcher in _run_celery_tasks_synchronously():
            self.addCleanup(patcher.stop)

    def _upload(self, rows, headers=HEADERS):
        return self.client.post(
            reverse("clients_web:child-import-upload"), {"file": _xlsx_file(rows, headers=headers)}
        )

    def _full_mapping_data(self, upload_response):
        data = {
            "headers_json": upload_response.context["headers_json"],
            "raw_rows_json": upload_response.context["raw_rows_json"],
            "meta_json": upload_response.context["meta_json"],
        }
        for key, current in _mapping_dict(upload_response.context["mapping_rows"]).items():
            if current:
                data[f"mapping_{key}"] = current
        return data

    def test_confirm_with_full_mapping_starts_a_dry_run_job(self):
        upload = self._upload(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )

        response = self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), self._full_mapping_data(upload)
        )

        job = ImportJob.objects.for_tenant(self.org).get()
        self.assertRedirects(
            response, reverse("clients_web:child-import-job-status", args=[job.id])
        )
        self.assertEqual(job.job_type, ImportJob.JobType.DRY_RUN)
        self.assertEqual(job.status, ImportJob.Status.DONE)
        self.assertEqual(job.ready_count, 1)
        self.assertEqual(job.error_count, 0)

    def test_confirm_reports_invalid_rows_as_errors_in_the_dry_run(self):
        upload = self._upload(
            [
                ["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"],
                ["", "не дата", "?", "", "не телефон", ""],
            ]
        )

        self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), self._full_mapping_data(upload)
        )

        job = ImportJob.objects.for_tenant(self.org).get()
        self.assertEqual(job.ready_count, 1)
        self.assertEqual(job.error_count, 1)

    def test_confirm_missing_required_field_rerenders_mapping_with_error(self):
        upload = self._upload(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )
        data = {
            "headers_json": upload.context["headers_json"],
            "raw_rows_json": upload.context["raw_rows_json"],
            "meta_json": upload.context["meta_json"],
            "mapping_child_name": "ФИО ребёнка",
        }

        response = self.client.post(reverse("clients_web:child-import-mapping-confirm"), data)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Телефон родителя", response.context["missing_required"])
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 0)
        self.assertFalse(ImportJob.objects.for_tenant(self.org).exists())

    def test_confirm_saves_mapping_for_reuse(self):
        upload = self._upload(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )

        self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), self._full_mapping_data(upload)
        )

        self.assertEqual(ImportColumnMapping.objects.for_tenant(self.org).count(), 1)

    def test_invalid_json_redirects_to_upload(self):
        response = self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), {"headers_json": "not json"}
        )

        self.assertRedirects(response, reverse("clients_web:child-import-upload"))

    def test_unknown_direction_is_a_warning_not_an_error(self):
        headers = [*HEADERS, "Направление"]
        upload = self._upload(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама", "Балет"]],
            headers=headers,
        )

        self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), self._full_mapping_data(upload)
        )

        job = ImportJob.objects.for_tenant(self.org).get()
        self.assertEqual(job.ready_count, 0)
        self.assertEqual(job.warning_count, 1)
        self.assertEqual(job.error_count, 0)
        self.assertTrue(any("направление" in w for w in job.report_rows[0]["messages"]))

    def test_no_child_or_parent_records_are_created_by_a_dry_run(self):
        # Критерий приёмки: сухой прогон не создаёт ни одной записи.
        upload = self._upload(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )

        self.client.post(
            reverse("clients_web:child-import-mapping-confirm"), self._full_mapping_data(upload)
        )

        self.assertEqual(Child.objects.for_tenant(self.org).count(), 0)


class ChildImportExecuteViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)
        for patcher in _run_celery_tasks_synchronously():
            self.addCleanup(patcher.stop)

    def _dry_run_job(self, rows, headers=HEADERS):
        upload = self.client.post(
            reverse("clients_web:child-import-upload"),
            {"file": _xlsx_file(rows, headers=headers)},
        )
        data = {
            "headers_json": upload.context["headers_json"],
            "raw_rows_json": upload.context["raw_rows_json"],
            "meta_json": upload.context["meta_json"],
        }
        for key, current in _mapping_dict(upload.context["mapping_rows"]).items():
            if current:
                data[f"mapping_{key}"] = current
        self.client.post(reverse("clients_web:child-import-mapping-confirm"), data)
        return ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.DRY_RUN)

    def test_execute_creates_an_execute_job_and_redirects_to_its_status(self):
        dry_run_job = self._dry_run_job(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )

        response = self.client.post(
            reverse("clients_web:child-import-execute", args=[dry_run_job.id])
        )

        execute_job = ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.EXECUTE)
        self.assertRedirects(
            response, reverse("clients_web:child-import-job-status", args=[execute_job.id])
        )
        self.assertEqual(execute_job.status, ImportJob.Status.DONE)
        self.assertEqual(execute_job.created_count, 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

    def test_execute_excludes_error_rows(self):
        dry_run_job = self._dry_run_job(
            [
                ["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"],
                ["", "не дата", "?", "", "не телефон", ""],
            ]
        )

        self.client.post(reverse("clients_web:child-import-execute", args=[dry_run_job.id]))

        execute_job = ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.EXECUTE)
        self.assertEqual(execute_job.total_rows, 1)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

    def test_repeated_execute_does_not_start_a_second_import(self):
        dry_run_job = self._dry_run_job(
            [["Данияр", "10.03.2018", "м", "Иванова Марина", "+77011234567", "мама"]]
        )
        url = reverse("clients_web:child-import-execute", args=[dry_run_job.id])

        self.client.post(url)
        response = self.client.post(url)

        execute_job = ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.EXECUTE)
        self.assertRedirects(
            response, reverse("clients_web:child-import-job-status", args=[execute_job.id])
        )
        dry_run_job.refresh_from_db()
        self.assertEqual(dry_run_job.executed_job, execute_job)
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

    def test_execute_on_pending_dry_run_job_returns_404(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            total_rows=0,
            rows_payload=[],
        )

        response = self.client.post(reverse("clients_web:child-import-execute", args=[job.id]))

        self.assertEqual(response.status_code, 404)

    def test_execute_on_another_organizations_job_returns_404(self):
        other_org = Organization.objects.create(name="Other Studio", slug="other-studio")
        job = ImportJob.objects.create(
            organization=other_org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            status=ImportJob.Status.DONE,
            total_rows=0,
            rows_payload=[],
        )

        response = self.client.post(reverse("clients_web:child-import-execute", args=[job.id]))

        self.assertEqual(response.status_code, 404)


class ChildImportJobStatusViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.other_org = Organization.objects.create(name="Other Studio", slug="other-studio")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)

    def test_pending_job_shows_in_progress_message(self):
        job = ImportJob.objects.create(
            organization=self.org, created_by=self.owner, total_rows=1, rows_payload=[]
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Ожидает")

    def test_done_execute_job_shows_result_counts(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.EXECUTE,
            total_rows=1,
            rows_payload=[],
            status=ImportJob.Status.DONE,
            created_count=1,
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertContains(response, "Готово")

    def test_done_dry_run_job_shows_report_and_execute_button(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            total_rows=2,
            rows_payload=[],
            status=ImportJob.Status.DONE,
            ready_count=1,
            warning_count=1,
            error_count=0,
            report_rows=[
                {"row_number": 2, "level": "ready", "child_name": "Данияр", "messages": []},
                {
                    "row_number": 3,
                    "level": "warning",
                    "child_name": "Айгерим",
                    "messages": ["дубликат — уже есть в базе"],
                },
            ],
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertContains(response, "дубликат")
        self.assertContains(response, reverse("clients_web:child-import-execute", args=[job.id]))
        self.assertContains(
            response, reverse("clients_web:child-import-report-download", args=[job.id])
        )

    def test_done_dry_run_job_with_no_ready_rows_hides_execute_button(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            total_rows=1,
            rows_payload=[],
            status=ImportJob.Status.DONE,
            ready_count=0,
            warning_count=0,
            error_count=1,
            report_rows=[
                {"row_number": 2, "level": "error", "child_name": "", "messages": ["ошибка"]}
            ],
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertNotContains(response, reverse("clients_web:child-import-execute", args=[job.id]))

    def test_already_executed_dry_run_job_links_to_import_instead_of_button(self):
        execute_job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.EXECUTE,
            total_rows=1,
            rows_payload=[],
        )
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            total_rows=1,
            rows_payload=[],
            status=ImportJob.Status.DONE,
            ready_count=1,
            executed_job=execute_job,
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertNotContains(response, reverse("clients_web:child-import-execute", args=[job.id]))
        self.assertContains(
            response, reverse("clients_web:child-import-job-status", args=[execute_job.id])
        )

    def test_failed_job_shows_error_message(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            total_rows=1,
            rows_payload=[],
            status=ImportJob.Status.FAILED,
            error_message="Что-то сломалось",
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertContains(response, "Что-то сломалось")

    def test_job_from_another_organization_is_not_visible(self):
        job = ImportJob.objects.create(
            organization=self.other_org, created_by=self.owner, total_rows=1, rows_payload=[]
        )

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertEqual(response.status_code, 404)


class ChildImportReportDownloadViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.other_org = Organization.objects.create(name="Other Studio", slug="other-studio")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)

    def _done_dry_run_job(self, organization=None):
        return ImportJob.objects.create(
            organization=organization or self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            status=ImportJob.Status.DONE,
            total_rows=1,
            rows_payload=[],
            report_rows=[
                {
                    "row_number": 147,
                    "level": "error",
                    "child_name": "",
                    "messages": ["не разобран телефон"],
                }
            ],
        )

    def test_download_returns_a_readable_xlsx_with_the_report(self):
        job = self._done_dry_run_job()

        response = self.client.get(
            reverse("clients_web:child-import-report-download", args=[job.id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0], ("Строка", "Статус", "ФИО ребёнка", "Сообщения"))
        self.assertEqual(rows[1][0], 147)
        self.assertIn("не разобран телефон", rows[1][3])

    def test_download_for_pending_job_returns_404(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.DRY_RUN,
            total_rows=1,
            rows_payload=[],
        )

        response = self.client.get(
            reverse("clients_web:child-import-report-download", args=[job.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_download_for_execute_job_returns_404(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            job_type=ImportJob.JobType.EXECUTE,
            status=ImportJob.Status.DONE,
            total_rows=1,
            rows_payload=[],
        )

        response = self.client.get(
            reverse("clients_web:child-import-report-download", args=[job.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_download_for_another_organizations_job_returns_404(self):
        job = self._done_dry_run_job(organization=self.other_org)

        response = self.client.get(
            reverse("clients_web:child-import-report-download", args=[job.id])
        )

        self.assertEqual(response.status_code, 404)


class ChildImportDuplicateDecisionsViewTests(TestCase):
    """Экран решений по дублям: по строке и массово для однотипных."""

    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)
        for patcher in _run_celery_tasks_synchronously():
            self.addCleanup(patcher.stop)
        # Существующая семья: два новых ребёнка этой мамы в файле — два
        # совпадения вида «родитель уже есть в базе».
        child = Child.objects.create(
            organization=self.org,
            full_name="Данияр",
            birth_date=datetime.date(2018, 3, 10),
            gender=Child.Gender.MALE,
        )
        parent = ParentContact.objects.create(organization=self.org, full_name="Иванова Марина")
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77011234567"
        )
        ChildContact.objects.create(
            organization=self.org, child=child, parent_contact=parent, role="mother"
        )

    def _dry_run_job(self):
        upload = self.client.post(
            reverse("clients_web:child-import-upload"),
            {
                "file": _xlsx_file(
                    [
                        ["Айгерим", "01.02.2017", "ж", "Иванова Марина", "+77011234567", "мама"],
                        ["Алия", "05.05.2019", "ж", "Иванова Марина", "+77011234567", "мама"],
                    ]
                )
            },
        )
        data = {
            "headers_json": upload.context["headers_json"],
            "raw_rows_json": upload.context["raw_rows_json"],
            "meta_json": upload.context["meta_json"],
        }
        for key, current in _mapping_dict(upload.context["mapping_rows"]).items():
            if current:
                data[f"mapping_{key}"] = current
        self.client.post(reverse("clients_web:child-import-mapping-confirm"), data)
        return ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.DRY_RUN)

    def test_status_page_lists_matches_with_decision_choices(self):
        job = self._dry_run_job()

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertEqual(len(response.context["duplicate_rows"]), 2)
        self.assertContains(response, "Привязать к существующему родителю")
        self.assertContains(response, 'name="decision_2"')
        self.assertEqual(response.context["duplicate_kinds"][0]["count"], 2)

    def test_bulk_decision_applies_to_all_rows_of_that_kind(self):
        job = self._dry_run_job()

        self.client.post(
            reverse("clients_web:child-import-decisions", args=[job.id]),
            {"bulk_kind": DuplicateKind.FAMILY, "bulk_decision": Decision.SKIP},
        )

        job.refresh_from_db()
        self.assertEqual(job.decisions, {"2": Decision.SKIP, "3": Decision.SKIP})

    def test_per_row_decision_is_saved_and_invalid_one_ignored(self):
        job = self._dry_run_job()

        self.client.post(
            reverse("clients_web:child-import-decisions", args=[job.id]),
            {"decision_2": Decision.CREATE_NEW, "decision_3": "удалить всё"},
        )

        job.refresh_from_db()
        self.assertEqual(job.decisions, {"2": Decision.CREATE_NEW})

    def test_execute_applies_decisions_submitted_with_the_form(self):
        job = self._dry_run_job()

        self.client.post(
            reverse("clients_web:child-import-execute", args=[job.id]),
            {"decision_2": Decision.SKIP, "decision_3": Decision.ATTACH},
        )

        execute_job = ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.EXECUTE)
        self.assertEqual(execute_job.skipped_count, 1)
        self.assertEqual(execute_job.attached_count, 1)
        self.assertEqual(ParentContact.objects.for_tenant(self.org).count(), 1)
        self.assertTrue(Child.objects.for_tenant(self.org).filter(full_name="Алия").exists())
        self.assertFalse(Child.objects.for_tenant(self.org).filter(full_name="Айгерим").exists())

    def test_decisions_cannot_change_after_import_started(self):
        job = self._dry_run_job()
        self.client.post(reverse("clients_web:child-import-execute", args=[job.id]))

        self.client.post(
            reverse("clients_web:child-import-decisions", args=[job.id]),
            {"bulk_kind": DuplicateKind.FAMILY, "bulk_decision": Decision.SKIP},
        )

        job.refresh_from_db()
        self.assertEqual(job.decisions, {})


class ChildImportRollbackViewTests(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.owner = User.objects.create_user(
            phone="+77010000001",
            full_name="Owner",
            password="pass12345",
            organization=self.org,
            role=User.Role.OWNER,
        )
        self.client.force_login(self.owner)
        for patcher in _run_celery_tasks_synchronously():
            self.addCleanup(patcher.stop)

    def _executed_job(self):
        upload = self.client.post(
            reverse("clients_web:child-import-upload"),
            {
                "file": _xlsx_file(
                    [["Данияр", "10.03.2018", "м", "Иванова", "+77011234567", "мама"]]
                )
            },
        )
        data = {
            "headers_json": upload.context["headers_json"],
            "raw_rows_json": upload.context["raw_rows_json"],
            "meta_json": upload.context["meta_json"],
        }
        for key, current in _mapping_dict(upload.context["mapping_rows"]).items():
            if current:
                data[f"mapping_{key}"] = current
        self.client.post(reverse("clients_web:child-import-mapping-confirm"), data)
        dry_run = ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.DRY_RUN)
        self.client.post(reverse("clients_web:child-import-execute", args=[dry_run.id]))
        return ImportJob.objects.for_tenant(self.org).get(job_type=ImportJob.JobType.EXECUTE)

    def test_done_import_shows_rollback_button_and_counts(self):
        job = self._executed_job()

        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertContains(response, reverse("clients_web:child-import-rollback", args=[job.id]))
        self.assertContains(response, "Создано родителей")

    def test_rollback_removes_imported_records(self):
        job = self._executed_job()
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)

        response = self.client.post(reverse("clients_web:child-import-rollback", args=[job.id]))

        self.assertRedirects(
            response, reverse("clients_web:child-import-job-status", args=[job.id])
        )
        self.assertEqual(Child.objects.for_tenant(self.org).count(), 0)
        job.refresh_from_db()
        self.assertIsNotNone(job.rolled_back_at)

    def test_rollback_blocked_after_work_started_keeps_data_and_explains_why(self):
        job = self._executed_job()
        child = Child.objects.for_tenant(self.org).get()
        CommunicationLog.objects.create(child=child, note="Позвонили", author=self.owner)

        self.client.post(reverse("clients_web:child-import-rollback", args=[job.id]))
        response = self.client.get(reverse("clients_web:child-import-job-status", args=[job.id]))

        self.assertEqual(Child.objects.for_tenant(self.org).count(), 1)
        self.assertNotContains(
            response, reverse("clients_web:child-import-rollback", args=[job.id])
        )
        self.assertContains(response, "записи коммуникаций")

    def test_running_job_shows_progress(self):
        job = ImportJob.objects.create(
            organization=self.org,
            created_by=self.owner,
            total_rows=5000,
            rows_payload=[],
            status=ImportJob.Status.RUNNING,
        )

        with mock.patch(
            "domains.people.clients.progress.get_progress",
            return_value={"phase": "writing", "done": 1200, "total": 5000},
        ):
            response = self.client.get(
                reverse("clients_web:child-import-job-status", args=[job.id])
            )

        self.assertContains(response, "1200 / 5000")
        self.assertContains(response, "location.reload")
