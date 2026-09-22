"""
Импорт детей из Excel (ТЗ п. 4.1, MVP критерий приёмки №1: "дубли
выявлены при импорте") — первый из трёх потребителей ChildService
(services.py, TRU-8 контракт №2). Не своя копия поиска дублей — весь
дедуп идёт через ChildService.find_duplicates(), иначе этот экран и
будущая конвертация заявки (M2) посчитают дубли по-разному.

Формат файла — фиксированные колонки (см. EXPECTED_HEADERS), без
произвольного сопоставления: администратору проще один раз привести файл
к формату, чем каждый раз сопоставлять колонки вручную, а сопоставление
"на глаз" само может стать источником ошибок импорта.

Дедуп внутри файла и второй ребёнок в семье внутри ОДНОГО файла — тот же
сценарий из ТЗ ("тот же телефон, другой ребёнок"), просто оба ребёнка
приходят в одном файле, а не по одному через веб-форму: строки
обрабатываются по порядку, и телефон, once встреченный, "запоминается"
на время предпросмотра/выполнения (см. _PhoneResolutionMap) — второй
ребёнок в файле с тем же номером не создаёт вторую маму.
"""

import datetime
import re
from dataclasses import dataclass, field

import openpyxl
from django.db.models import Q

from domains.platform.core.phone import InvalidPhoneNumberError, normalize_phone_number

from .models import Child, ChildContact, ParentContact
from .services import ChildService, DuplicateReason


def _existing_family_child_keys(organization, phone):
    """(имя.lower(), дата рождения) уже существующих детей всех
    ParentContact с этим телефоном — используется, чтобы засеять
    _PhoneResolutionMap реальным состоянием базы в момент, когда телефон
    впервые встретился в файле: без этого второй ряд файла с тем же
    телефоном сверялся бы только с детьми из ЭТОГО файла, забывая, что
    у найденной семьи уже есть свои дети в базе (см. resolve_rows)."""
    parent_ids = ParentContact.objects.for_tenant(organization).filter(
        Q(phones__number=phone) | Q(whatsapp=phone)
    )
    children = Child.objects.for_tenant(organization).filter(
        contacts__parent_contact__in=parent_ids
    )
    return {(c.full_name.lower(), c.birth_date) for c in children}


# Порядок значим только для сообщения об ошибке заголовка — сами колонки
# ищутся по названию, не по позиции (администратор может переставить
# столбцы местами в своей копии файла).
EXPECTED_HEADERS = {
    "ФИО ребёнка": "child_name",
    "Дата рождения ребёнка": "birth_date",
    "Пол ребёнка": "gender",
    "ФИО родителя": "parent_name",
    "Телефон родителя": "phone",
    "Роль родителя": "role",
    "Медицинские заметки": "medical_notes",
    "Остаток занятий": "reported_balance",
}
OPTIONAL_HEADERS = {"Роль родителя", "Медицинские заметки", "Остаток занятий"}
REQUIRED_HEADERS = [h for h in EXPECTED_HEADERS if h not in OPTIONAL_HEADERS]

GENDER_ALIASES = {
    "м": Child.Gender.MALE,
    "мужской": Child.Gender.MALE,
    "male": Child.Gender.MALE,
    "ж": Child.Gender.FEMALE,
    "женский": Child.Gender.FEMALE,
    "female": Child.Gender.FEMALE,
}
ROLE_ALIASES = {label.lower(): value for value, label in ChildContact.Role.choices}
ROLE_ALIASES.update(
    {
        "мама": ChildContact.Role.MOTHER,
        "папа": ChildContact.Role.FATHER,
        "бабушка": ChildContact.Role.GRANDMOTHER,
        "опекун": ChildContact.Role.GUARDIAN,
    }
)

# "мама Айгерим" вместо чистого ФИО в колонке родителя — реальная грязь
# из ТЗ (найдено на синтетических файлах, см. docs/import_format.md).
# Роль-префикс отделяется от имени и используется, только если отдельная
# колонка "Роль родителя" не задана явно (она приоритетнее).
_ROLE_PREFIX_RE = re.compile(
    r"^("
    + "|".join(re.escape(k) for k in sorted(ROLE_ALIASES, key=len, reverse=True))
    + r")\s+(.+)$",
    re.IGNORECASE,
)


def _split_role_prefix(raw_name: str) -> tuple[str | None, str]:
    match = _ROLE_PREFIX_RE.match(raw_name.strip())
    if not match:
        return None, raw_name
    prefix, rest = match.groups()
    return ROLE_ALIASES.get(prefix.lower()), rest.strip()


# Несколько телефонов в одной ячейке через запятую/точку с запятой/слэш/
# "и" — тоже реальная грязь (родитель и его партнёр в одной колонке).
# ContactPhone — отдельная модель (не одно поле), поэтому все валидные
# номера сохраняются, не только первый.
_PHONE_SPLIT_RE = re.compile(r"[,;/]+|\s+и\s+", re.IGNORECASE)


def _split_phones(raw: str) -> list[str]:
    parts = [p.strip() for p in _PHONE_SPLIT_RE.split(raw) if p.strip()]
    return parts or ([raw.strip()] if raw.strip() else [])


class RowAction:
    CREATE_NEW_FAMILY = "create_new_family"
    ATTACH_EXISTING = "attach_existing"
    SKIP = "skip"
    ERROR = "error"


@dataclass
class ImportRow:
    row_number: int  # номер строки в файле (для сообщений администратору)
    child_name: str = ""
    birth_date: datetime.date | None = None
    gender: str = ""
    parent_name: str = ""
    phone: str = ""  # первый валидный номер — используется для дедупа
    extra_phones: list[str] = field(
        default_factory=list
    )  # остальные валидные номера из той же ячейки
    role: str = ChildContact.Role.OTHER
    medical_notes: str = ""
    # Сырое значение колонки "Остаток занятий" — не превращается в
    # реальный Subscription при импорте (решение зафиксировано в
    # docs/import_format.md: перенос остатков — отдельная задача домена
    # Bekzat'а), но и не отбрасывается молча — видно в отчёте импорта,
    # чтобы не всплыло на приёмке.
    reported_balance: str = ""
    errors: list[str] = field(default_factory=list)

    # Заполняется resolve_rows() — не на этапе разбора файла.
    action: str = RowAction.CREATE_NEW_FAMILY
    reason: str | None = None  # DuplicateReason.* — почему предложено именно это действие
    matched_child: Child | None = None
    matched_parent_id: str | None = None
    # Объект (не только id) — только для отрисовки предпросмотра, не
    # сериализуется в rows_json (после подтверждения ChildService сам
    # найдёт родителя по matched_parent_id).
    matched_parent: object | None = None
    # Номер строки файла, чья семья "усыновит" этого ребёнка — заполнен,
    # когда matched_parent_id ещё не существует в базе (родителя создаст
    # ЭТА строка при выполнении, см. execute_import): предпросмотру нечего
    # показать как id, но показать "строка N" — можно и нужно.
    attached_to_row_number: int | None = None

    @property
    def is_valid(self) -> bool:
        return not self.errors

    @property
    def role_label(self) -> str:
        try:
            return ChildContact.Role(self.role).label
        except ValueError:
            return self.role

    @property
    def gender_label(self) -> str:
        try:
            return Child.Gender(self.gender).label
        except ValueError:
            return self.gender


# Три формата дат из реальной практики (ТЗ): точки, слэши, ISO — плюс
# нативная дата Excel (обрабатывается отдельно ниже, это не строка).
_DATE_FORMATS = ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d")


def _parse_cell_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if isinstance(value, str) and value.strip():
        for fmt in _DATE_FORMATS:
            try:
                return datetime.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _clean_str(value) -> str:
    return str(value).strip() if value is not None else ""


def _build_merge_lookup(sheet) -> dict[tuple[int, int], object]:
    """(строка, колонка, 1-based) -> значение якоря объединённой ячейки —
    объединённые ячейки — реальная грязь (ТЗ): например, ФИО/телефон
    родителя объединены на несколько строк его детей, и без этой подсказки
    все строки кроме первой читались бы как пустые."""
    lookup = {}
    for merged_range in sheet.merged_cells.ranges:
        anchor = sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                if (row, col) != (merged_range.min_row, merged_range.min_col):
                    lookup[(row, col)] = anchor
    return lookup


def parse_workbook(file) -> tuple[list[ImportRow], list[str]]:
    """
    Возвращает (строки, ошибки_заголовка). При ошибке заголовка строки
    пустые — нет смысла разбирать данные под неверными колонками.

    read_only=False (не потоковый режим) — нужен доступ к
    merged_cells.ranges и произвольным ячейкам для _build_merge_lookup();
    для разового admin-действия на файл в тысячи строк это не проблема
    производительности (не сравнимо с бюджетом списка на 5000 детей).
    """
    workbook = openpyxl.load_workbook(file, data_only=True, read_only=False)
    sheet = workbook.active
    merge_lookup = _build_merge_lookup(sheet)

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header_index = {
        _clean_str(cell): idx for idx, cell in enumerate(header_row) if _clean_str(cell)
    }
    missing = [h for h in REQUIRED_HEADERS if h not in header_index]
    if missing:
        return [], [f"В файле не найдены обязательные колонки: {', '.join(missing)}"]

    rows = []
    for row_number, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(_clean_str(v) for v in raw_row):
            continue  # пустая строка — обычный хвост файла, не ошибка

        def cell(header, _raw_row=raw_row, _row_number=row_number):
            idx = header_index.get(header)
            if idx is None:
                return None
            value = _raw_row[idx] if idx < len(_raw_row) else None
            if value is None:
                value = merge_lookup.get((_row_number, idx + 1))
            return value

        row = ImportRow(row_number=row_number)
        errors = []

        row.child_name = _clean_str(cell("ФИО ребёнка"))
        if not row.child_name:
            errors.append("не заполнено ФИО ребёнка")

        birth_date = _parse_cell_date(cell("Дата рождения ребёнка"))
        if birth_date is None:
            errors.append("не удалось разобрать дату рождения (ожидается ДД.МM.ГГГГ)")
        row.birth_date = birth_date

        gender_raw = _clean_str(cell("Пол ребёнка")).lower()
        gender = GENDER_ALIASES.get(gender_raw)
        if gender is None:
            errors.append(f"не удалось разобрать пол ребёнка: {gender_raw!r}")
        row.gender = gender or ""

        parent_name_raw = _clean_str(cell("ФИО родителя"))
        prefix_role, cleaned_parent_name = _split_role_prefix(parent_name_raw)
        row.parent_name = cleaned_parent_name
        if not row.parent_name:
            errors.append("не заполнено ФИО родителя")

        phone_raw = _clean_str(cell("Телефон родителя"))
        normalized_phones = []
        for candidate in _split_phones(phone_raw):
            try:
                normalized = normalize_phone_number(candidate)
            except InvalidPhoneNumberError:
                continue
            if normalized not in normalized_phones:
                normalized_phones.append(normalized)
        if not normalized_phones:
            errors.append(f"не похоже на телефон: {phone_raw!r}")
            row.phone = phone_raw
        else:
            row.phone = normalized_phones[0]
            row.extra_phones = normalized_phones[1:]

        # Явная колонка "Роль родителя" приоритетнее роли, угаданной из
        # префикса в имени ("мама Айгерим") — администратор мог заполнить
        # обе, и явная колонка — более осознанный ввод.
        role_raw = _clean_str(cell("Роль родителя")).lower()
        row.role = ROLE_ALIASES.get(role_raw) or prefix_role or ChildContact.Role.OTHER

        row.medical_notes = _clean_str(cell("Медицинские заметки"))
        row.reported_balance = _clean_str(cell("Остаток занятий"))

        row.errors = errors
        rows.append(row)

    return rows, []


class _PhoneResolutionMap:
    """Телефон -> уже найденный/созданный родитель в рамках ОДНОГО прогона
    (предпросмотра или выполнения) — второй ребёнок с тем же номером
    внутри файла попадает к тому же родителю, а не заводит второго."""

    def __init__(self):
        self._by_phone: dict[str, dict] = {}

    def get(self, phone):
        return self._by_phone.get(phone)

    def remember(self, phone, *, parent_id, row_number, child_keys=()):
        entry = self._by_phone.setdefault(
            phone, {"parent_id": parent_id, "row_number": row_number, "children": set()}
        )
        entry["parent_id"] = parent_id
        entry["row_number"] = row_number
        entry["children"].update(child_keys)

    @staticmethod
    def child_key(child_name, birth_date):
        return (child_name.lower(), birth_date)


def resolve_rows(organization, rows: list[ImportRow]) -> None:
    """Проставляет action/reason/matched_* на каждой валидной строке —
    дедуп и внутри файла, и против базы (см. докстринг модуля). Строки с
    ошибками разбора не резолвятся вообще (action не имеет смысла)."""
    phone_map = _PhoneResolutionMap()

    for row in rows:
        if not row.is_valid:
            row.action = RowAction.ERROR
            continue

        remembered = phone_map.get(row.phone)
        if remembered:
            child_key = _PhoneResolutionMap.child_key(row.child_name, row.birth_date)
            if child_key in remembered["children"]:
                # Тот же телефон И тот же ребёнок, что уже встречались в
                # этом файле — повтор строки, не второй ребёнок.
                row.action = RowAction.SKIP
                row.reason = DuplicateReason.PHONE
            else:
                row.action = RowAction.ATTACH_EXISTING
                row.reason = DuplicateReason.EXISTING_PARENT_NEW_CHILD
                row.matched_parent_id = remembered["parent_id"]
                if remembered["parent_id"] is None:
                    row.attached_to_row_number = remembered["row_number"]
                phone_map.remember(
                    row.phone,
                    parent_id=remembered["parent_id"],
                    row_number=remembered["row_number"],
                    child_keys=[child_key],
                )
            continue

        matches = ChildService.find_duplicates(
            organization,
            phone=row.phone,
            child_name=row.child_name,
            birth_date=row.birth_date,
        )
        phone_match = next((m for m in matches if m.reason == DuplicateReason.PHONE), None)
        family_match = next(
            (m for m in matches if m.reason == DuplicateReason.EXISTING_PARENT_NEW_CHILD), None
        )
        weak_match = next(
            (
                m
                for m in matches
                if m.reason in (DuplicateReason.NAME_AND_BIRTH_DATE, DuplicateReason.NAME_ONLY)
            ),
            None,
        )

        if phone_match:
            # Тот же телефон И тот же ребёнок — похоже на повтор строки/
            # повторную загрузку того же файла. По умолчанию пропускаем,
            # администратор может явно выбрать "создать всё равно".
            row.action = RowAction.SKIP
            row.reason = DuplicateReason.PHONE
            row.matched_child = phone_match.child
        elif family_match:
            row.action = RowAction.ATTACH_EXISTING
            row.reason = DuplicateReason.EXISTING_PARENT_NEW_CHILD
            row.matched_parent_id = str(family_match.parent.id)
            row.matched_parent = family_match.parent
            existing_keys = _existing_family_child_keys(organization, row.phone)
            existing_keys.add(_PhoneResolutionMap.child_key(row.child_name, row.birth_date))
            phone_map.remember(
                row.phone,
                parent_id=row.matched_parent_id,
                row_number=row.row_number,
                child_keys=existing_keys,
            )
        else:
            row.action = RowAction.CREATE_NEW_FAMILY
            if weak_match:
                # Слабый сигнал — показать, но не настаивать (ТЗ): не
                # блокирует и не меняет действие по умолчанию.
                row.reason = weak_match.reason
                row.matched_child = weak_match.child
            phone_map.remember(
                row.phone,
                parent_id=None,
                row_number=row.row_number,
                child_keys=[_PhoneResolutionMap.child_key(row.child_name, row.birth_date)],
            )


@dataclass
class ImportResult:
    created: int = 0
    attached_to_existing_family: int = 0
    skipped: int = 0
    failed: list[tuple[int, str]] = field(default_factory=list)  # (row_number, error)
    # Строки, где была непустая колонка "Остаток занятий" — сознательно
    # НЕ превращается в реальный Subscription при импорте (см.
    # docs/import_format.md), но и не пропадает молча: администратор
    # должен увидеть и перенести остаток вручную/отдельным механизмом.
    unhandled_balances: list[tuple[int, str, str]] = field(
        default_factory=list
    )  # (row_number, child_name, reported_balance)


def execute_import(organization, rows: list[ImportRow]) -> ImportResult:
    """Строки — по порядку файла (важно для второго ребёнка в семье в
    рамках одного прогона, см. _PhoneResolutionMap)."""
    result = ImportResult()
    phone_map = _PhoneResolutionMap()

    for row in rows:
        if row.reported_balance and row.action not in (RowAction.SKIP, RowAction.ERROR):
            result.unhandled_balances.append((row.row_number, row.child_name, row.reported_balance))

        if row.action in (RowAction.SKIP, RowAction.ERROR):
            result.skipped += 1
            continue

        try:
            child_data = {
                "full_name": row.child_name,
                "birth_date": row.birth_date,
                "gender": row.gender,
                "medical_notes": row.medical_notes,
            }

            remembered = phone_map.get(row.phone)
            if remembered:
                parent_data = {"id": remembered["parent_id"]}
            elif row.action == RowAction.ATTACH_EXISTING and row.matched_parent_id:
                parent_data = {"id": row.matched_parent_id}
            else:
                parent_data = {
                    "full_name": row.parent_name,
                    "phones": [row.phone, *row.extra_phones],
                }

            child = ChildService.create_with_parent(
                organization, child_data=child_data, parent_data=parent_data, link_role=row.role
            )
            parent_id = str(ChildContact.objects.get(child=child).parent_contact_id)
            phone_map.remember(row.phone, parent_id=parent_id, row_number=row.row_number)

            if row.action == RowAction.ATTACH_EXISTING:
                result.attached_to_existing_family += 1
            else:
                result.created += 1
        except Exception as exc:  # noqa: BLE001 — одна плохая строка не должна ронять весь импорт
            result.failed.append((row.row_number, str(exc)))

    return result
