from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.db import connection
from django.test import tag
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APITestCase

from domains.money.payments.services import record_payment
from domains.people.clients.models import Child, ChildContact, ContactPhone, ParentContact
from domains.platform.core.utils import today_for_org
from domains.platform.tenants.models import Branch, Direction, Organization
from domains.platform.users.models import User

from .debt import debt_by_child
from .models import Subscription
from .subscription_types import create_type


class DebtorsApiTests(APITestCase):
    """TRU-68: экран «Задолженности» — те же суммы, что список детей и карточка
    родителя; давность по дате центра; порог просрочки из настроек."""

    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.abay = Branch.objects.create(organization=self.org, name="Абая")
        self.orbita = Branch.objects.create(organization=self.org, name="Орбита")
        self.ballet = Direction.objects.create(organization=self.org, name="Балет")
        self.type = create_type(
            self.org,
            name="8 занятий",
            price=30000,
            quota_sessions=8,
            duration_days=30,
            directions=[self.ballet],
        )
        self.admin = User.objects.create_user(
            phone="+77010000001",
            password="x",
            full_name="Админ",
            organization=self.org,
            role=User.Role.ADMIN,
        )
        self.today = today_for_org(self.org)
        self.fresh = self.sub("Свежий Долг", days_ago=1, paid=10000)  # долг 20000
        self.old = self.sub("Старый Долг", days_ago=40, paid=25000, branch=self.orbita)  # 5000
        self.sub("Всё Оплачено", days_ago=10, paid=30000)
        # Переплата по второму абонементу не гасит долг по первому (как в списке детей).
        self.sub("Старый Долг", days_ago=3, paid=35000, child=self.old.child)
        parent = ParentContact.objects.create(
            organization=self.org, full_name="Мама Свежего", whatsapp="+77071112233"
        )
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77071112233"
        )
        ChildContact.objects.create(
            organization=self.org, child=self.fresh.child, parent_contact=parent, is_payer=True
        )

    def sub(self, name, *, days_ago, paid, branch=None, child=None):
        child = child or Child.objects.create(
            organization=self.org, full_name=name, birth_date=self.today - timedelta(days=3000)
        )
        subscription = Subscription.objects.create(
            organization=self.org,
            child=child,
            subscription_type_version=self.type.versions.latest(),
            direction=self.ballet,
            branch=branch or self.abay,
            starts_on=self.today - timedelta(days=days_ago),
            ends_on=self.today + timedelta(days=30 - days_ago),
            list_price=30000,
            price=30000,
        )
        if paid:
            record_payment(actor=self.admin, subscription=subscription, amount=paid, method="cash")
        return subscription

    def get(self, **params):
        self.client.force_authenticate(self.admin)
        return self.client.get("/api/v1/subscriptions/debtors/", params)

    def test_list_total_and_same_sums_as_children_list(self):
        data = self.get().data
        self.assertEqual(data["count"], 2)
        self.assertEqual(data["total_debt"], "25000")
        self.assertEqual((data["overdue_count"], data["overdue_debt"]), (1, "5000"))
        self.assertEqual([r["child_name"] for r in data["results"]], ["Свежий Долг", "Старый Долг"])
        by_child = debt_by_child(self.org, [self.fresh.child_id, self.old.child_id])
        for row in data["results"]:
            child = Child.objects.get(pk=row["child_id"])
            self.assertEqual(Decimal(row["debt"]), by_child[child.pk])

    def test_overdue_uses_org_threshold(self):
        data = self.get(overdue="1").data
        self.assertEqual(data["overdue_days"], 5)
        self.assertEqual([r["child_name"] for r in data["results"]], ["Старый Долг"])
        self.org.settings = {"debt_overdue_days_threshold": 60}
        self.org.save()
        self.assertEqual(self.get(overdue="1").data["count"], 0)

    def test_sort_by_age_oldest_first(self):
        data = self.get(sort="age", dir="desc").data
        self.assertEqual([r["age_days"] for r in data["results"]], [40, 1])
        self.assertTrue(data["results"][0]["overdue"])

    def test_sort_by_age_newest_first(self):
        # TRU-133: «по возрастанию давности» — сначала самые свежие долги.
        data = self.get(sort="age", dir="asc").data
        self.assertEqual([r["age_days"] for r in data["results"]], [1, 40])

    def test_filters_branch_and_search(self):
        self.assertEqual(self.get(branch=str(self.orbita.pk)).data["count"], 1)
        self.assertEqual(self.get(q="свеж").data["results"][0]["child_name"], "Свежий Долг")

    def test_payer_contact_and_reminder(self):
        row = self.get(q="Свежий").data["results"][0]
        self.assertEqual(row["parent_name"], "Мама Свежего")
        self.assertEqual(row["whatsapp"], "+77071112233")
        self.assertIn("20 000 ₸", row["reminder_text"])

    def test_whatsapp_falls_back_to_first_phone(self):
        # TRU-133: whatsapp не заполнен — кнопка WhatsApp всё равно есть, на первый телефон.
        parent = ParentContact.objects.create(organization=self.org, full_name="Папа Старого")
        ContactPhone.objects.create(
            organization=self.org, parent_contact=parent, number="+77075556677"
        )
        ChildContact.objects.create(
            organization=self.org, child=self.old.child, parent_contact=parent, is_payer=True
        )
        row = self.get(q="Старый").data["results"][0]
        self.assertEqual(row["whatsapp"], "+77075556677")
        self.assertIn("5 000 ₸", row["reminder_text"])

    def test_page_queries_do_not_grow_with_rows(self):
        # TRU-133: плательщик и телефоны — пачкой на страницу, не запросом на строку.
        # Свежий пользователь на каждый запрос: allowed_branch_ids кэширует
        # филиалы на объекте, со вторым запросом запросов стало бы меньше.
        url = "/api/v1/subscriptions/debtors/"
        self.client.force_authenticate(User.objects.get(pk=self.admin.pk))
        with CaptureQueriesContext(connection) as before:
            self.client.get(url)
        for i in range(5):
            sub = self.sub(f"Ещё Долг {i}", days_ago=2, paid=0)
            parent = ParentContact.objects.create(organization=self.org, full_name=f"Родитель {i}")
            ContactPhone.objects.create(
                organization=self.org, parent_contact=parent, number=f"+7707000000{i}"
            )
            ChildContact.objects.create(
                organization=self.org, child=sub.child, parent_contact=parent, is_payer=True
            )
        self.client.force_authenticate(User.objects.get(pk=self.admin.pk))
        with CaptureQueriesContext(connection) as after:
            data = self.client.get(url).data
        self.assertEqual(data["count"], 7)
        self.assertEqual(len(after.captured_queries), len(before.captured_queries))

    def test_phones_hidden_without_permission(self):
        accountant = User.objects.create_user(
            phone="+77010000002",
            password="x",
            full_name="Бухгалтер",
            organization=self.org,
            role=User.Role.ACCOUNTANT,
        )
        self.client.force_authenticate(accountant)
        row = self.client.get("/api/v1/subscriptions/debtors/", {"q": "Свежий"}).data["results"][0]
        if not accountant_can_view_phone(accountant):
            self.assertIsNone(row["phone"])
            self.assertIsNone(row["whatsapp"])

    def test_teacher_has_no_access(self):
        teacher = User.objects.create_user(
            phone="+77010000003",
            password="x",
            full_name="Педагог",
            organization=self.org,
            role=User.Role.TEACHER,
        )
        self.client.force_authenticate(teacher)
        self.assertEqual(self.client.get("/api/v1/subscriptions/debtors/").status_code, 403)

    def test_export_opens_in_excel(self):
        self.client.force_authenticate(self.admin)
        response = self.client.get("/api/v1/subscriptions/debtors/export/")
        self.assertEqual(response.status_code, 200)
        book = openpyxl.load_workbook(io_bytes(response.content))
        sheet = book.active
        self.assertEqual(sheet.cell(row=2, column=1).value, "Свежий Долг")
        self.assertEqual(sheet.cell(row=sheet.max_row, column=1).value, "Итого")


@tag("tenant_isolation")
class DebtorsFilterIsolationTests(APITestCase):
    """TRU-133: ?branch=/?direction=/?group= ищутся только в своей организации.
    Чужой или кривой id — пустой список, а не «все филиалы» и не 500."""

    def setUp(self):
        self.org = Organization.objects.create(name="True Ballet", slug="true-ballet")
        self.other = Organization.objects.create(name="Другой центр", slug="other")
        self.branch = Branch.objects.create(organization=self.org, name="Абая")
        self.direction = Direction.objects.create(organization=self.org, name="Балет")
        self.foreign_branch = Branch.objects.create(organization=self.other, name="Чужой")
        self.foreign_direction = Direction.objects.create(organization=self.other, name="Чужое")
        type_ = create_type(
            self.org,
            name="8 занятий",
            price=30000,
            quota_sessions=8,
            duration_days=30,
            directions=[self.direction],
        )
        today = today_for_org(self.org)
        child = Child.objects.create(
            organization=self.org, full_name="Должник", birth_date=today - timedelta(days=3000)
        )
        Subscription.objects.create(
            organization=self.org,
            child=child,
            subscription_type_version=type_.versions.latest(),
            direction=self.direction,
            branch=self.branch,
            starts_on=today - timedelta(days=28),
            ends_on=today + timedelta(days=2),
            list_price=30000,
            price=30000,
        )
        self.admin = User.objects.create_user(
            phone="+77010000009",
            password="x",
            full_name="Админ",
            organization=self.org,
            role=User.Role.ADMIN,
        )
        self.client.force_authenticate(self.admin)

    def test_debtors_own_filters_work(self):
        response = self.client.get(
            "/api/v1/subscriptions/debtors/",
            {"branch": str(self.branch.pk), "direction": str(self.direction.pk)},
        )
        self.assertEqual(response.data["count"], 1)

    def test_debtors_foreign_or_broken_filter_gives_empty_list(self):
        for params in (
            {"branch": str(self.foreign_branch.pk)},
            {"direction": str(self.foreign_direction.pk)},
            {"direction": "not-a-uuid"},
        ):
            with self.subTest(params=params):
                response = self.client.get("/api/v1/subscriptions/debtors/", params)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data["count"], 0)
                self.assertEqual(response.data["total_debt"], "0")
                response = self.client.get("/api/v1/subscriptions/debtors/export/", params)
                self.assertEqual(response.status_code, 200)

    def test_renewals_foreign_filter_gives_empty_list(self):
        self.assertEqual(self.client.get("/api/v1/subscriptions/renewals/").data["count"], 1)
        for params in (
            {"branch": str(self.foreign_branch.pk)},
            {"direction": str(self.foreign_direction.pk)},
            {"group": "not-a-uuid"},
        ):
            with self.subTest(params=params):
                response = self.client.get("/api/v1/subscriptions/renewals/", params)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.data["count"], 0)


def accountant_can_view_phone(user):
    from domains.platform.core.role_permissions import can_view_phone

    return can_view_phone(user)


def io_bytes(content):
    import io

    return io.BytesIO(content)
